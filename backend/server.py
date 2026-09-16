"""KLINIK SUEL MOTOR - Backend API (FastAPI + MongoDB)."""
import io
import logging
import os
import re
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Annotated, Any, Literal, Optional

import jwt
import requests
from dotenv import load_dotenv
from fastapi import APIRouter, Depends, FastAPI, File, HTTPException, UploadFile
from pymongo import InsertOne, UpdateOne
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import RedirectResponse, Response, StreamingResponse
from fastapi.security import OAuth2PasswordBearer
from jwt.exceptions import InvalidTokenError
from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import Workbook, load_workbook
from pwdlib import PasswordHash
from pydantic import BaseModel, Field
from starlette.middleware.cors import CORSMiddleware

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")
logger = logging.getLogger("suel")

client = AsyncIOMotorClient(os.environ["MONGO_URL"])
db = client[os.environ["DB_NAME"]]

JWT_SECRET = os.environ["JWT_SECRET"]
JWT_EXPIRE_MINUTES = int(os.environ["JWT_EXPIRE_MINUTES"])
ALGORITHM = "HS256"
Role = Literal["owner", "kasir", "mekanik", "partman"]
ALL_ROLES: tuple[str, ...] = ("owner", "kasir", "mekanik", "partman")

password_hash = PasswordHash.recommended()
DUMMY_HASH = password_hash.hash("dummy-password-never-used")
oauth2 = OAuth2PasswordBearer(tokenUrl="/api/auth/login")

DEFAULT_SHOP = {
    "id": "shop",
    "name": "KLINIK SUEL MOTOR",
    "address": "Jl. Puyung-Bonjeruk, Dusun Batu Ngereng, Desa Bunkate, Kec. Jonggat, Kab. Lombok Tengah, NTB 83561",
    "phone": "081917106463",
    "whatsapp": "081917106463",
    "postal_code": "83561",
    "logo_path": None,
    "logo_url": "https://customer-assets-4nw71qhi.emergentagent.net/job_motor-service-hub-13/artifacts/qwui9q90_WhatsApp%20Image%202026-06-27%20at%2013.50.22.jpeg",
    "logo_version": 1,
}


async def get_shop() -> dict:
    doc = await db.settings.find_one({"id": "shop"}, {"_id": 0})
    return {**DEFAULT_SHOP, **(doc or {})}


def wa_phone(raw: str) -> str:
    phone = "".join(ch for ch in (raw or "") if ch.isdigit())
    if phone.startswith("0"):
        phone = "62" + phone[1:]
    return phone


async def verify_owner_password(pw: Optional[str]) -> bool:
    if not pw:
        return False
    async for u in db.users.find({"role": "owner", "disabled": False, "deleted_at": None}):
        if password_hash.verify(pw, u["hashed_password"]):
            return True
    return False


async def notify(roles: list[str], kind: str, title: str, body: str, trx_id: Optional[str] = None, meta: Optional[dict] = None, wa: bool = False):
    from urllib.parse import quote
    wa_url = None
    if wa:
        shop = await get_shop()
        phone = wa_phone(shop.get("whatsapp") or shop.get("phone"))
        wa_url = f"https://wa.me/{phone}?text={quote(title + chr(10) + body)}" if phone else None
    await db.notifications.insert_one({"id": new_id(), "roles": roles, "kind": kind, "title": title, "body": body, "transaction_id": trx_id, "meta": meta or {},
                                       "read": False, "read_by": [], "wa_url": wa_url, "created_at": iso(now())})


async def notify_owner(kind: str, title: str, body: str, trx_id: Optional[str] = None, meta: Optional[dict] = None):
    await notify(["owner"], kind, title, body, trx_id, meta, wa=True)

# --------------------------------------------------------------------------- helpers

def now() -> datetime:
    return datetime.now(timezone.utc)


def iso(dt: datetime) -> str:
    return dt.isoformat()


def new_id() -> str:
    return str(uuid.uuid4())


def clean(doc: Optional[dict]) -> Optional[dict]:
    if doc is None:
        return None
    doc.pop("_id", None)
    return doc


async def next_counter(key: str) -> int:
    doc = await db.counters.find_one_and_update(
        {"_id": key}, {"$inc": {"seq": 1}}, upsert=True, return_document=True
    )
    return int(doc["seq"])


async def log_status(trx_id: str, status: str, user: dict, note: str = ""):
    await db.transaction_status_logs.insert_one(
        {"id": new_id(), "transaction_id": trx_id, "status": status, "note": note,
         "user_id": user["id"], "username": user["username"], "created_at": iso(now())}
    )


async def audit(trx_id: Optional[str], action: str, user: dict, detail: str = "", meta: Optional[dict] = None):
    await db.audit_logs.insert_one(
        {"id": new_id(), "transaction_id": trx_id, "action": action, "detail": detail, "meta": meta or {},
         "user_id": user["id"], "username": user["username"], "created_at": iso(now())}
    )


async def set_status(trx_id: str, status: str, user: dict, note: str = "", extra: Optional[dict] = None):
    upd = {"status": status, "updated_at": iso(now())}
    if extra:
        upd.update(extra)
    await db.service_transactions.update_one({"id": trx_id}, {"$set": upd})
    await log_status(trx_id, status, user, note)
    await audit(trx_id, "STATUS", user, f"Status → {status}. {note}".strip())


# --------------------------------------------------------------------------- auth

class LoginIn(BaseModel):
    username: str
    password: str


class UserPublic(BaseModel):
    id: str
    username: str
    name: str
    role: Role
    disabled: bool = False


class UserCreate(BaseModel):
    username: str
    password: str
    name: str
    role: Role


class UserUpdate(BaseModel):
    name: Optional[str] = None
    role: Optional[Role] = None
    password: Optional[str] = None
    disabled: Optional[bool] = None


def make_token(user: dict) -> str:
    n = now()
    return jwt.encode(
        {"sub": user["username"], "role": user["role"], "iat": n, "exp": n + timedelta(minutes=JWT_EXPIRE_MINUTES)},
        JWT_SECRET, algorithm=ALGORITHM,
    )


def to_public(doc: dict) -> dict:
    return {"id": doc["id"], "username": doc["username"], "name": doc.get("name", doc["username"]),
            "role": doc["role"], "disabled": doc.get("disabled", False)}


async def current_user(token: Annotated[str, Depends(oauth2)]) -> dict:
    err = HTTPException(status_code=401, detail="Sesi tidak valid atau sudah habis", headers={"WWW-Authenticate": "Bearer"})
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[ALGORITHM])
        username = payload.get("sub")
        if not isinstance(username, str):
            raise err
    except InvalidTokenError:
        raise err
    doc = await db.users.find_one({"username": username, "deleted_at": None})
    if not doc or doc.get("disabled"):
        raise err
    return to_public(doc)


def require_roles(*allowed: str):
    async def dep(user: Annotated[dict, Depends(current_user)]) -> dict:
        if user["role"] != "owner" and user["role"] not in allowed:
            raise HTTPException(status_code=403, detail="Anda tidak memiliki akses untuk aksi ini")
        return user
    return dep


CurrentUser = Annotated[dict, Depends(current_user)]
OwnerUser = Annotated[dict, Depends(require_roles())]
# Mekanik dapat mengerjakan semua tugas kasir & partman (akses penuh operasional).
KasirUser = Annotated[dict, Depends(require_roles("kasir", "mekanik"))]
MekanikUser = Annotated[dict, Depends(require_roles("mekanik"))]
PartmanUser = Annotated[dict, Depends(require_roles("partman", "mekanik"))]
StockUser = Annotated[dict, Depends(require_roles("partman", "mekanik"))]
SalesUser = Annotated[dict, Depends(require_roles("partman", "kasir", "mekanik"))]
RegistrarUser = Annotated[dict, Depends(require_roles("mekanik", "kasir"))]

# --------------------------------------------------------------------------- app

@asynccontextmanager
async def lifespan(_: FastAPI):
    await db.users.create_index("username", unique=True)
    await db.customers.create_index("phone")
    await db.vehicles.create_index("plate")
    await db.parts.create_index("code")
    await db.service_transactions.create_index([("created_at", -1)])
    await db.service_items.create_index("transaction_id")
    await seed()
    yield
    client.close()


async def seed():
    defaults = [("owner", "owner123", "Pemilik Bengkel", "owner"), ("kasir", "kasir123", "Kasir", "kasir"),
                ("mekanik", "mekanik123", "Mekanik Andi", "mekanik"), ("partman", "partman123", "Partman", "partman")]
    for username, pw, name, role in defaults:
        await db.users.update_one(
            {"username": username},
            {"$setOnInsert": {"id": new_id(), "username": username, "name": name, "role": role,
                              "hashed_password": password_hash.hash(pw), "disabled": False, "created_at": iso(now())}},
            upsert=True,
        )
    if await db.services.count_documents({}) == 0:
        for code, name, price in [("J001", "Servis Ringan", 50000), ("J002", "Servis CVT", 50000), ("J003", "Ganti Oli (jasa)", 15000),
                                  ("J004", "Ganti Kampas Rem", 35000), ("J005", "Servis Besar", 150000), ("J006", "Ganti Bearing", 40000)]:
            await db.services.insert_one({"id": new_id(), "code": code, "name": name, "price": price, "active": True, "created_at": iso(now()), "deleted_at": None})
    if await db.parts.count_documents({}) == 0:
        for code, name, price, cost, stock, barcode in [
            ("P001", "Oli Mesin MPX2 0.8L", 55000, 45000, 20, "8990001000011"), ("P002", "Kampas Rem Depan", 45000, 32000, 8, "8990001000028"),
            ("P003", "Bearing 6002", 25000, 15000, 5, "8990001000035"), ("P004", "Busi NGK CPR8EA", 20000, 14000, 12, "8990001000042"),
            ("P005", "Filter Udara Beat", 40000, 28000, 3, "8990001000059"), ("P006", "Roller CVT Set", 60000, 42000, 6, "8990001000066")]:
            await db.parts.insert_one({"id": new_id(), "code": code, "name": name, "barcode": barcode, "price": price, "cost": cost,
                                       "stock": stock, "min_stock": 5, "unit": "pcs", "active": True, "created_at": iso(now()), "deleted_at": None})
    # SOP checklist default (owner boleh ubah)
    if await db.service_checklist_items.count_documents({}) == 0:
        for i, name in enumerate(["Baut Kencang", "Oli Cek", "Test Jalan", "Tools Lengkap"]):
            await db.service_checklist_items.insert_one({"id": new_id(), "name": name, "seq": i, "active": True, "created_at": iso(now())})
    # Pengaturan keuangan default (owner isi saldo awal sendiri)
    await db.settings.update_one({"id": "finance"}, {"$setOnInsert": {"id": "finance", "saldo_awal_bengkel": 0, "saldo_awal_pribadi": 0, "owner_draw": 0}}, upsert=True)
    # Hak akses per role default (semua fitur aktif). owner selalu penuh.
    await db.settings.update_one({"id": "access"}, {"$setOnInsert": {"id": "access", "roles": {}}}, upsert=True)
    # Data karyawan default dari user seed
    if await db.employees.count_documents({}) == 0:
        for name, pos, base, bonus in [("Mekanik Andi", "mekanik", 1500000, 10000), ("Kasir", "kasir", 1200000, 0), ("Partman", "partman", 1200000, 0)]:
            await db.employees.insert_one({"id": new_id(), "name": name, "position": pos, "base_salary": base,
                                           "bonus_per_unit": bonus, "active": True, "created_at": iso(now()), "deleted_at": None})
    # --- Migrasi data lama (idempoten) ---
    # 1) Tandai pembayaran milik transaksi DIBATALKAN agar tidak mencemari laporan
    cancelled_ids = [t["id"] async for t in db.service_transactions.find({"status": "DIBATALKAN"}, {"id": 1})]
    if cancelled_ids:
        await db.payments.update_many({"transaction_id": {"$in": cancelled_ids}}, {"$set": {"cancelled": True}})
    # 2) Bersihkan cancel_reason yang berisi password (mis. 'owner123')
    await db.service_transactions.update_many(
        {"cancel_reason": {"$regex": r"^\s*(owner|admin|kasir|mekanik|partman)?\d*123\s*$", "$options": "i"}},
        {"$set": {"cancel_reason": "(alasan tidak dicatat)"}})


app = FastAPI(title="Klinik Suel Motor API", lifespan=lifespan)
api = APIRouter(prefix="/api")


@api.get("/")
async def root():
    return {"app": "KLINIK SUEL MOTOR", "ok": True}


@api.post("/auth/login")
async def login(body: LoginIn):
    doc = await db.users.find_one({"username": body.username.strip().lower()})
    if not doc:
        password_hash.verify(body.password, DUMMY_HASH)
        raise HTTPException(status_code=401, detail="Username atau password salah")
    if doc.get("disabled") or not password_hash.verify(body.password, doc["hashed_password"]):
        raise HTTPException(status_code=401, detail="Username atau password salah")
    return {"access_token": make_token(doc), "token_type": "bearer", "user": to_public(doc)}


@api.get("/auth/me")
async def me(user: CurrentUser):
    return user


# --------------------------------------------------------------------------- users (owner)

@api.get("/users")
async def list_users(_: OwnerUser):
    docs = await db.users.find({"deleted_at": None}, {"_id": 0, "hashed_password": 0}).to_list(500)
    return docs


@api.delete("/users/{uid}")
async def delete_user(uid: str, user: OwnerUser):
    if uid == user["id"]:
        raise HTTPException(status_code=400, detail="Tidak dapat menghapus akun sendiri")
    doc = await db.users.find_one_and_update({"id": uid, "deleted_at": None}, {"$set": {"deleted_at": iso(now()), "disabled": True}}, return_document=True)
    if not doc:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    await audit(None, "USER_DELETE", user, f"Hapus pengguna {doc['username']}")
    return {"ok": True}


@api.post("/users")
async def create_user(body: UserCreate, user: OwnerUser):
    uname = body.username.strip().lower()
    if await db.users.find_one({"username": uname}):
        raise HTTPException(status_code=400, detail="Username sudah dipakai")
    doc = {"id": new_id(), "username": uname, "name": body.name, "role": body.role,
           "hashed_password": password_hash.hash(body.password), "disabled": False, "created_at": iso(now())}
    await db.users.insert_one(doc)
    await audit(None, "USER_CREATE", user, f"Tambah pengguna {uname} ({body.role})")
    return to_public(doc)


@api.put("/users/{uid}")
async def update_user(uid: str, body: UserUpdate, user: OwnerUser):
    upd: dict[str, Any] = {k: v for k, v in body.model_dump(exclude_none=True).items() if k != "password"}
    if body.password:
        upd["hashed_password"] = password_hash.hash(body.password)
    doc = await db.users.find_one_and_update({"id": uid}, {"$set": upd}, return_document=True)
    if not doc:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    await audit(None, "USER_UPDATE", user, f"Ubah pengguna {doc['username']}")
    return to_public(doc)


@api.get("/mechanics")
async def list_mechanics(_: CurrentUser):
    docs = await db.users.find({"role": {"$in": ["mekanik", "owner"]}, "disabled": False, "deleted_at": None}, {"_id": 0, "hashed_password": 0}).to_list(200)
    return docs


# --------------------------------------------------------------------------- customers & vehicles

class CustomerIn(BaseModel):
    name: str
    phone: str = ""
    address: str = ""
    notes: str = ""
    dusun: str = ""
    desa: str = ""
    kecamatan: str = ""
    kabupaten: str = ""


class VehicleIn(BaseModel):
    customer_id: str
    plate: str
    brand: str = ""
    model: str = ""
    year: str = ""
    color: str = ""
    chassis_no: str = ""
    engine_no: str = ""
    km_last: int = 0


def regex(q: str) -> dict:
    return {"$regex": q.strip(), "$options": "i"}


@api.get("/customers")
async def list_customers(_: CurrentUser, q: str = ""):
    flt: dict[str, Any] = {"deleted_at": None}
    if q.strip():
        plates = await db.vehicles.find({"plate": regex(q), "deleted_at": None}, {"customer_id": 1}).to_list(50)
        flt["$or"] = [{"name": regex(q)}, {"phone": regex(q)}, {"id": {"$in": [p["customer_id"] for p in plates]}}]
    docs = await db.customers.find(flt, {"_id": 0}).sort("name", 1).to_list(200)
    return docs


@api.post("/customers")
async def create_customer(body: CustomerIn, user: CurrentUser):
    seq = await next_counter("customer")
    doc = {"id": new_id(), "code": f"C{seq:04d}", **body.model_dump(), "created_at": iso(now()), "deleted_at": None}
    await db.customers.insert_one(doc)
    await audit(None, "CUSTOMER_CREATE", user, f"Pelanggan baru {body.name}")
    return clean(doc)


@api.put("/customers/{cid}")
async def update_customer(cid: str, body: CustomerIn, user: CurrentUser):
    doc = await db.customers.find_one_and_update({"id": cid}, {"$set": body.model_dump()}, return_document=True)
    if not doc:
        raise HTTPException(status_code=404, detail="Pelanggan tidak ditemukan")
    await audit(None, "CUSTOMER_UPDATE", user, f"Ubah pelanggan {body.name}")
    return clean(doc)


@api.get("/vehicles")
async def list_vehicles(_: CurrentUser, q: str = "", customer_id: str = ""):
    flt: dict[str, Any] = {"deleted_at": None}
    if customer_id:
        flt["customer_id"] = customer_id
    if q.strip():
        flt["$or"] = [{"plate": regex(q)}, {"brand": regex(q)}, {"model": regex(q)}]
    docs = await db.vehicles.find(flt, {"_id": 0}).sort("plate", 1).to_list(300)
    cust_ids = list({d["customer_id"] for d in docs})
    custs = {c["id"]: c for c in await db.customers.find({"id": {"$in": cust_ids}}, {"_id": 0}).to_list(300)}
    for d in docs:
        c = custs.get(d["customer_id"])
        d["customer_name"] = c["name"] if c else "-"
    return docs


@api.post("/vehicles")
async def create_vehicle(body: VehicleIn, user: CurrentUser):
    plate = body.plate.strip().upper()
    if await db.vehicles.find_one({"plate": plate, "deleted_at": None}):
        raise HTTPException(status_code=400, detail="Nomor polisi sudah terdaftar")
    doc = {"id": new_id(), **body.model_dump(), "plate": plate, "last_service_at": None, "created_at": iso(now()), "deleted_at": None}
    await db.vehicles.insert_one(doc)
    await audit(None, "VEHICLE_CREATE", user, f"Motor baru {plate}")
    return clean(doc)


@api.put("/vehicles/{vid}")
async def update_vehicle(vid: str, body: VehicleIn, user: CurrentUser):
    doc = await db.vehicles.find_one_and_update({"id": vid}, {"$set": {**body.model_dump(), "plate": body.plate.strip().upper()}}, return_document=True)
    if not doc:
        raise HTTPException(status_code=404, detail="Motor tidak ditemukan")
    await audit(None, "VEHICLE_UPDATE", user, f"Ubah motor {doc['plate']}")
    return clean(doc)


# --------------------------------------------------------------------------- master jasa & parts

class ServiceIn(BaseModel):
    code: str
    name: str
    price: int
    active: bool = True


class PartIn(BaseModel):
    code: str
    name: str
    barcode: str = ""
    price: int
    cost: int = 0
    stock: int = 0
    min_stock: int = 5
    unit: str = "pcs"
    rack: str = ""
    active: bool = True


class StockAdjustIn(BaseModel):
    qty: int
    reason: str = "PENYESUAIAN"
    note: str = ""
    date: str = ""  # owner boleh mundur-tanggalkan (DD/MM/YYYY / YYYY-MM-DD)


@api.get("/services")
async def list_services(_: CurrentUser, q: str = ""):
    flt: dict[str, Any] = {"deleted_at": None}
    if q.strip():
        flt["$or"] = [{"code": regex(q)}, {"name": regex(q)}]
    return await db.services.find(flt, {"_id": 0}).sort("name", 1).to_list(500)


@api.post("/services")
async def create_service(body: ServiceIn, user: OwnerUser):
    doc = {"id": new_id(), **body.model_dump(), "created_at": iso(now()), "deleted_at": None}
    await db.services.insert_one(doc)
    await audit(None, "SERVICE_CREATE", user, f"Jasa baru {body.name}")
    return clean(doc)


@api.put("/services/{sid}")
async def update_service(sid: str, body: ServiceIn, user: OwnerUser):
    doc = await db.services.find_one_and_update({"id": sid}, {"$set": body.model_dump()}, return_document=True)
    if not doc:
        raise HTTPException(status_code=404, detail="Jasa tidak ditemukan")
    await audit(None, "SERVICE_UPDATE", user, f"Ubah jasa {body.name}")
    return clean(doc)


@api.get("/parts")
async def list_parts(_: CurrentUser, q: str = "", barcode: str = "", low: bool = False):
    flt: dict[str, Any] = {"deleted_at": None}
    if barcode.strip():
        flt["$or"] = [{"barcode": barcode.strip()}, {"code": barcode.strip()}]
    elif q.strip():
        flt["$or"] = [{"code": regex(q)}, {"name": regex(q)}, {"barcode": regex(q)}, {"rack": regex(q)}]
    if low:
        flt["$expr"] = {"$lte": ["$stock", "$min_stock"]}
    docs = await db.parts.find(flt, {"_id": 0}).sort("name", 1).to_list(None)
    for d in docs:
        d["low_stock"] = int(d.get("stock", 0) or 0) <= int(d.get("min_stock", 0) or 0)
    return docs


@api.post("/parts")
async def create_part(body: PartIn, user: StockUser):
    code = body.code.strip()
    if is_junk_part(code, body.name):
        await db.import_trash.insert_one({"id": new_id(), "code": code, "name": body.name, "barcode": body.barcode,
                                          "price": body.price, "cost": body.cost, "stock": body.stock, "min_stock": body.min_stock,
                                          "unit": body.unit, "rack": body.rack, "reason": "Part sampah (kode Z/L atau nama Kosong/Low)",
                                          "source": "manual", "created_by": user["username"], "created_at": iso(now())})
        raise HTTPException(status_code=400, detail="Part terdeteksi sebagai part sampah (kode Z/L atau nama Kosong/Low) dan dipindahkan ke Sampah Import, bukan ke master.")
    if await db.parts.find_one({"code": code, "deleted_at": None}):
        raise HTTPException(status_code=400, detail="Kode part sudah ada")
    doc = {"id": new_id(), **body.model_dump(), "code": code, "created_at": iso(now()), "deleted_at": None}
    await db.parts.insert_one(doc)
    if body.stock:
        await db.stock_movements.insert_one({"id": new_id(), "part_id": doc["id"], "part_code": doc["code"], "part_name": doc["name"],
                                             "qty": body.stock, "stock_before": 0, "stock_after": body.stock, "reason": "STOK AWAL",
                                             "transaction_id": None, "transaction_no": None, "username": user["username"], "created_at": iso(now())})
    await audit(None, "PART_CREATE", user, f"Part baru {body.name}")
    return clean(doc)


@api.put("/parts/{pid}")
async def update_part(pid: str, body: PartIn, user: StockUser):
    old = await db.parts.find_one({"id": pid})
    if not old:
        raise HTTPException(status_code=404, detail="Part tidak ditemukan")
    upd = body.model_dump()
    upd.pop("stock", None)  # stock only changes through adjust endpoint
    doc = await db.parts.find_one_and_update({"id": pid}, {"$set": upd}, return_document=True)
    await audit(None, "PART_UPDATE", user, f"Ubah part {body.name}", {"before_price": old.get("price"), "after_price": body.price})
    return clean(doc)


@api.post("/parts/{pid}/adjust")
async def adjust_stock(pid: str, body: StockAdjustIn, user: StockUser):
    part = await db.parts.find_one({"id": pid})
    if not part:
        raise HTTPException(status_code=404, detail="Part tidak ditemukan")
    before = int(part.get("stock", 0))
    after = before + body.qty
    if after < 0:
        raise HTTPException(status_code=400, detail="Stok tidak boleh negatif")
    created = iso(now())
    if body.date and user["role"] == "owner":
        d = parse_date(body.date)
        if d:
            created = datetime(d.year, d.month, d.day, 12, 0, tzinfo=timezone(timedelta(hours=7))).astimezone(timezone.utc).isoformat()
    doc = await db.parts.find_one_and_update({"id": pid}, {"$set": {"stock": after}}, return_document=True)
    await db.stock_movements.insert_one({"id": new_id(), "part_id": pid, "part_code": part["code"], "part_name": part["name"], "qty": body.qty,
                                         "stock_before": before, "stock_after": after, "reason": body.reason, "note": body.note,
                                         "transaction_id": None, "transaction_no": None, "username": user["username"], "created_at": created})
    min_stock = int(part.get("min_stock", 0) or 0)
    if body.qty < 0 and after <= min_stock < before:
        await notify(["owner", "partman"], "STOK MENIPIS", f"STOK {'HABIS' if after <= 0 else 'MENIPIS'}: {part['name']}",
                     f"Kode {part['code']} · sisa {after} {part.get('unit', 'pcs')} (minimum {min_stock}). Segera lakukan pembelian.", None, {"part_id": pid})
    await audit(None, "STOCK_ADJUST", user, f"{part['name']} {body.qty:+d} ({body.reason})")
    return clean(doc)


@api.get("/parts/{pid}/movements")
async def part_movements(pid: str, _: CurrentUser):
    return await db.stock_movements.find({"part_id": pid}, {"_id": 0}).sort("created_at", -1).to_list(200)


@api.get("/stock-movements")
async def all_movements(_: CurrentUser):
    return await db.stock_movements.find({}, {"_id": 0}).sort("created_at", -1).to_list(300)


@api.put("/stock-movements/{mid}/date")
async def edit_movement_date(mid: str, body: dict, user: OwnerUser):
    mv = await db.stock_movements.find_one({"id": mid})
    if not mv:
        raise HTTPException(status_code=404, detail="Mutasi tidak ditemukan")
    d = parse_date(str(body.get("date", "")))
    if not d:
        raise HTTPException(status_code=400, detail="Tanggal tidak valid")
    wib = timezone(timedelta(hours=7))
    try:
        old = datetime.fromisoformat(mv["created_at"]).astimezone(wib)
    except (ValueError, TypeError, KeyError):
        old = datetime(d.year, d.month, d.day, 12, 0, tzinfo=wib)
    new_dt = old.replace(year=d.year, month=d.month, day=d.day)
    await db.stock_movements.update_one({"id": mid}, {"$set": {"created_at": new_dt.astimezone(timezone.utc).isoformat()}})
    await audit(None, "MOVEMENT_DATE_EDIT", user, f"Ubah tanggal mutasi {mv.get('part_name')} → {d.strftime('%d/%m/%Y')}")
    return {"ok": True}


# --------------------------------------------------------------------------- cek fisik stok (stock opname)

def rack_sort_key(rack: str):
    """Urutan alami kode rak: A1.01 < A1.2 < A2 < B1. Part tanpa rak di urutan akhir."""
    s = (rack or "").strip()
    if not s:
        return (1, [])
    return (0, [(1, int(t)) if t.isdigit() else (0, t) for t in re.findall(r"\d+|\D+", s.lower())])


async def get_stock_check(cid: str) -> dict:
    chk = await db.stock_checks.find_one({"id": cid})
    if not chk:
        raise HTTPException(status_code=404, detail="Sesi cek fisik tidak ditemukan")
    return chk


def wib_date_str() -> str:
    return now().astimezone(timezone(timedelta(hours=7))).strftime("%Y-%m-%d")


@api.post("/stock-checks")
async def start_stock_check(user: PartmanUser):
    today = wib_date_str()
    existing = await db.stock_checks.find_one({"date": today, "status": "PROSES"}, {"_id": 0})
    if existing:
        return existing
    parts = await db.parts.find({"deleted_at": None}, {"_id": 0}).to_list(None)
    parts.sort(key=lambda p: rack_sort_key(p.get("rack")))
    chk = {"id": new_id(), "date": today, "status": "PROSES", "created_by": user["username"], "created_at": iso(now()),
           "submitted_at": None, "decided_at": None, "decided_by": None,
           "total_items": len(parts), "checked_count": 0, "diff_count": 0}
    await db.stock_checks.insert_one(chk)
    items = [{"id": new_id(), "check_id": chk["id"], "seq": i, "part_id": p["id"], "part_code": p["code"], "part_name": p["name"],
              "rack": p.get("rack") or "", "unit": p.get("unit", "pcs"), "stock_system": int(p.get("stock", 0) or 0),
              "stock_physical": None, "checked": False, "diff": None, "adj_status": None, "created_at": chk["created_at"]}
             for i, p in enumerate(parts)]
    if items:
        await db.stock_check_items.insert_many(items)
    chk.pop("_id", None)
    return chk


@api.get("/stock-checks")
async def list_stock_checks(_: PartmanUser):
    return await db.stock_checks.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)


@api.get("/stock-checks/{cid}")
async def stock_check_detail(cid: str, _: PartmanUser):
    chk = await get_stock_check(cid)
    items = await db.stock_check_items.find({"check_id": cid}, {"_id": 0}).sort("seq", 1).to_list(None)
    chk.pop("_id", None)
    return {**chk, "items": items}


class CheckItemIn(BaseModel):
    physical: int = 0
    checked: bool = True


@api.put("/stock-checks/{cid}/items/{iid}")
async def update_check_item(cid: str, iid: str, body: CheckItemIn, user: PartmanUser):
    chk = await get_stock_check(cid)
    if chk["status"] != "PROSES":
        raise HTTPException(status_code=400, detail="Sesi ini sudah tidak bisa diubah")
    if body.physical < 0:
        raise HTTPException(status_code=400, detail="Jumlah fisik tidak boleh negatif")
    item = await db.stock_check_items.find_one({"id": iid, "check_id": cid})
    if not item:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    physical = body.physical if body.checked else None
    diff = (physical - item["stock_system"]) if physical is not None else None
    await db.stock_check_items.update_one({"id": iid}, {"$set": {"stock_physical": physical, "checked": body.checked, "diff": diff}})
    checked = await db.stock_check_items.count_documents({"check_id": cid, "checked": True})
    diffs = await db.stock_check_items.count_documents({"check_id": cid, "checked": True, "diff": {"$ne": 0}})
    await db.stock_checks.update_one({"id": cid}, {"$set": {"checked_count": checked, "diff_count": diffs}})
    return clean(await db.stock_check_items.find_one({"id": iid}))


@api.post("/stock-checks/{cid}/submit")
async def submit_stock_check(cid: str, user: PartmanUser):
    chk = await get_stock_check(cid)
    if chk["status"] != "PROSES":
        raise HTTPException(status_code=400, detail="Sesi ini sudah diajukan / selesai")
    diffs = await db.stock_check_items.find({"check_id": cid, "checked": True, "diff": {"$ne": 0}}, {"_id": 0}).to_list(None)
    if not diffs:
        raise HTTPException(status_code=400, detail="Tidak ada selisih yang perlu diajukan")
    await db.stock_check_items.update_many({"check_id": cid, "checked": True, "diff": {"$ne": 0}}, {"$set": {"adj_status": "PENDING"}})
    await db.stock_checks.update_one({"id": cid}, {"$set": {"status": "DIAJUKAN", "submitted_at": iso(now()), "diff_count": len(diffs)}})
    lines = "\n".join(f"• {d['part_name']}: sistem {d['stock_system']} → fisik {d['stock_physical']} ({d['diff']:+d})" for d in diffs[:10])
    if len(diffs) > 10:
        lines += f"\n… dan {len(diffs) - 10} part lainnya"
    await notify(["owner"], "CEK FISIK STOK", f"PERSETUJUAN PENYESUAIAN STOK: {len(diffs)} part selisih",
                 f"@{user['username']} mengajukan penyesuaian stok dari cek fisik {chk['date']}:\n{lines}", None, {"check_id": cid})
    await audit(None, "STOCK_CHECK_SUBMIT", user, f"Cek fisik {chk['date']}: {len(diffs)} selisih diajukan")
    return clean(await db.stock_checks.find_one({"id": cid}))


@api.post("/stock-checks/{cid}/approve")
async def approve_stock_check(cid: str, user: OwnerUser):
    chk = await get_stock_check(cid)
    if chk["status"] != "DIAJUKAN":
        raise HTTPException(status_code=400, detail="Tidak ada pengajuan yang menunggu persetujuan")
    pendings = await db.stock_check_items.find({"check_id": cid, "adj_status": "PENDING"}, {"_id": 0}).to_list(None)
    now_iso = iso(now())
    applied = 0
    for it in pendings:
        part = await db.parts.find_one({"id": it["part_id"]})
        if not part or part.get("deleted_at"):
            continue
        before = int(part.get("stock", 0) or 0)
        after = int(it["stock_physical"])
        if before == after:
            continue
        await db.parts.update_one({"id": part["id"]}, {"$set": {"stock": after}})
        await db.stock_movements.insert_one({"id": new_id(), "part_id": part["id"], "part_code": part["code"], "part_name": part["name"],
                                             "qty": after - before, "stock_before": before, "stock_after": after,
                                             "reason": "CEK FISIK STOK", "note": f"Opname {chk['date']} oleh @{chk['created_by']}, disetujui @{user['username']}",
                                             "transaction_id": None, "transaction_no": None, "username": user["username"], "created_at": now_iso})
        min_stock = int(part.get("min_stock", 0) or 0)
        if after - before < 0 and after <= min_stock < before:
            await notify(["owner", "partman"], "STOK MENIPIS", f"STOK {'HABIS' if after <= 0 else 'MENIPIS'}: {part['name']}",
                         f"Kode {part['code']} · sisa {after} {part.get('unit', 'pcs')} (minimum {min_stock}). Segera lakukan pembelian.", None, {"part_id": part["id"]})
        applied += 1
    await db.stock_check_items.update_many({"check_id": cid, "adj_status": "PENDING"}, {"$set": {"adj_status": "DISETUJUI"}})
    await db.stock_checks.update_one({"id": cid}, {"$set": {"status": "DISETUJUI", "decided_at": now_iso, "decided_by": user["username"]}})
    await notify(["partman"], "CEK FISIK STOK", "PENYESUAIAN STOK DISETUJUI",
                 f"Owner menyetujui {len(pendings)} penyesuaian stok dari cek fisik {chk['date']}. Stok sistem telah diperbarui.", None, {"check_id": cid})
    await audit(None, "STOCK_CHECK_APPROVE", user, f"Cek fisik {chk['date']}: {applied} penyesuaian diterapkan")
    return clean(await db.stock_checks.find_one({"id": cid}))


@api.post("/stock-checks/{cid}/reject")
async def reject_stock_check(cid: str, user: OwnerUser):
    chk = await get_stock_check(cid)
    if chk["status"] != "DIAJUKAN":
        raise HTTPException(status_code=400, detail="Tidak ada pengajuan yang menunggu persetujuan")
    await db.stock_check_items.update_many({"check_id": cid, "adj_status": "PENDING"}, {"$set": {"adj_status": "DITOLAK"}})
    await db.stock_checks.update_one({"id": cid}, {"$set": {"status": "DITOLAK", "decided_at": iso(now()), "decided_by": user["username"]}})
    await notify(["partman"], "CEK FISIK STOK", "PENYESUAIAN STOK DITOLAK",
                 f"Owner menolak penyesuaian stok dari cek fisik {chk['date']}. Stok sistem tidak berubah.", None, {"check_id": cid})
    await audit(None, "STOCK_CHECK_REJECT", user, f"Cek fisik {chk['date']}: pengajuan ditolak")
    return clean(await db.stock_checks.find_one({"id": cid}))


@api.post("/stock-checks/{cid}/finish")
async def finish_stock_check(cid: str, user: PartmanUser):
    chk = await get_stock_check(cid)
    if chk["status"] != "PROSES":
        raise HTTPException(status_code=400, detail="Sesi ini sudah selesai")
    await db.stock_checks.update_one({"id": cid}, {"$set": {"status": "SELESAI"}})
    await audit(None, "STOCK_CHECK_FINISH", user, f"Cek fisik {chk['date']} selesai tanpa penyesuaian")
    return clean(await db.stock_checks.find_one({"id": cid}))


def parse_date(s: str):
    """Accepts DD/MM/YYYY, YYYY-MM-DD or ISO datetime. Returns date or None."""
    s = (s or "").strip()
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s).astimezone(timezone(timedelta(hours=7))).date()
    except ValueError:
        return None


def wib_today() -> str:
    return now().astimezone(timezone(timedelta(hours=7))).strftime("%Y%m%d")


# --------------------------------------------------------------------------- transactions

STATUS_ACTIVE = ["TERDAFTAR", "MENUNGGU_SERVIS", "DIPROSES", "MENUNGGU_PERSETUJUAN", "MENUNGGU_KASIR", "MENUNGGU_PEMBAYARAN"]
STATUS_PAID = ["SUDAH_DIBAYAR", "NOTA_DICETAK", "NOTA_TERKIRIM", "SELESAI"]


class ItemIn(BaseModel):
    kind: Literal["jasa", "part"]
    ref_id: Optional[str] = None  # service id or part id
    name: str
    code: str = ""
    price: int
    qty: int = 1
    note: str = ""
    owner_password: Optional[str] = None


class ItemEditIn(BaseModel):
    name: Optional[str] = None
    price: Optional[int] = None
    qty: Optional[int] = None
    note: Optional[str] = None
    reason: str = ""
    owner_password: Optional[str] = None


class TransactionCreate(BaseModel):
    customer_id: str
    vehicle_id: str
    km_in: int = 0
    complaint_main: str
    complaint_extra: str = ""
    initial_check: str = ""
    mechanic_id: Optional[str] = None
    items: list[ItemIn] = []
    date: str = ""  # owner boleh mundur-tanggalkan pendaftaran


class StartIn(BaseModel):
    mechanic_id: Optional[str] = None
    km: Optional[int] = None


class WorkNoteIn(BaseModel):
    work_done: str = ""
    inspection: str = ""
    mechanic_note: str = ""


class FinishIn(BaseModel):
    work_done: str = ""
    condition: str = ""
    mechanic_note: str = ""
    next_recommendation: str = ""
    next_km: Optional[int] = None
    next_date: str = ""


class PaymentIn(BaseModel):
    method: Literal["CASH", "TRANSFER", "QRIS", "DEBIT", "KREDIT"]
    amount_paid: int
    discount: int = 0
    reference: str = ""
    note: str = ""
    due_date: str = ""


class DebtPayIn(BaseModel):
    amount: int
    method: Literal["CASH", "TRANSFER", "QRIS", "DEBIT"] = "CASH"
    reference: str = ""
    note: str = ""


class PriceChangeIn(BaseModel):
    price: int
    reason: str


class CancelIn(BaseModel):
    reason: str = ""
    owner_password: Optional[str] = None


class OwnerPasswordIn(BaseModel):
    owner_password: str


async def apply_stock_delta(part_id: str, delta: int, reason: str, trx: Optional[dict], user: dict):
    """Change part stock by delta (negative = out) and record a movement."""
    if delta == 0:
        return
    p = await db.parts.find_one_and_update({"id": part_id}, {"$inc": {"stock": delta}}, return_document=True)
    if p:
        await db.stock_movements.insert_one({"id": new_id(), "part_id": p["id"], "part_code": p["code"], "part_name": p["name"], "qty": delta,
                                             "stock_before": p["stock"] - delta, "stock_after": p["stock"], "reason": reason,
                                             "transaction_id": trx["id"] if trx else None, "transaction_no": trx["trx_no"] if trx else None,
                                             "username": user["username"], "created_at": iso(now())})
        min_stock = int(p.get("min_stock", 0) or 0)
        if delta < 0 and p["stock"] <= min_stock < p["stock"] - delta:
            await notify(["owner", "partman"], "STOK MENIPIS", f"STOK {'HABIS' if p['stock'] <= 0 else 'MENIPIS'}: {p['name']}",
                         f"Kode {p['code']} · sisa {p['stock']} {p.get('unit', 'pcs')} (minimum {min_stock}). Segera lakukan pembelian.", None, {"part_id": p["id"]})


async def resync_paid(tid: str):
    """After editing a paid invoice, keep payment/invoice totals and debt in sync."""
    d = await build_detail(await get_trx(tid))
    total = d["totals"]["total"]
    pay = d.get("payment")
    if pay:
        paid = int(pay.get("amount_paid", 0))
        debt = max(total - paid, 0)
        change = max(paid - total, 0) if pay["method"] == "CASH" else 0
        await db.payments.update_one({"transaction_id": tid}, {"$set": {"total": total, "change": change}})
        await db.invoices.update_one({"transaction_id": tid}, {"$set": {"total": total}})
        await db.service_transactions.update_one({"id": tid}, {"$set": {"debt_amount": debt, "debt_status": "LUNAS" if debt == 0 else "BELUM LUNAS"}})
        await db.service_history.update_one({"transaction_id": tid}, {"$set": {"total": total}})


def item_doc(trx: dict, it: ItemIn, source: str, user: dict, approval: str, cost: int = 0) -> dict:
    # cost_snapshot: harga beli part saat transaksi dibuat -> untuk hitung laba (price - cost) * qty
    return {"id": new_id(), "transaction_id": trx["id"], "kind": it.kind, "ref_id": it.ref_id, "code": it.code, "name": it.name,
            "price": it.price, "cost": int(cost or 0), "qty": it.qty, "subtotal": it.price * it.qty,
            "profit": (it.price - int(cost or 0)) * it.qty if it.kind == "part" else it.price * it.qty,
            "note": it.note, "source": source,
            "approval": approval, "deleted_at": None, "added_by": user["username"], "created_at": iso(now())}


async def item_cost(it: ItemIn) -> int:
    """Ambil snapshot harga beli (cost) part dari master saat item dibuat."""
    if it.kind == "part" and it.ref_id:
        p = await db.parts.find_one({"id": it.ref_id}, {"_id": 0, "cost": 1})
        if p:
            return int(p.get("cost", 0) or 0)
    return 0


async def build_detail(trx: dict) -> dict:
    trx = clean(trx) or {}
    items = await db.service_items.find({"transaction_id": trx["id"], "deleted_at": None}, {"_id": 0}).sort("created_at", 1).to_list(500)
    part_ids = [i["ref_id"] for i in items if i["kind"] == "part" and i.get("ref_id")]
    parts = {p["id"]: p for p in await db.parts.find({"id": {"$in": part_ids}}, {"_id": 0}).to_list(500)}
    for i in items:
        if i["kind"] == "part" and i.get("ref_id") in parts:
            i["stock_available"] = parts[i["ref_id"]]["stock"]
            i["stock_ok"] = parts[i["ref_id"]]["stock"] >= i["qty"]
            i["rack"] = parts[i["ref_id"]].get("rack", "")
    counted = [i for i in items if i["approval"] == "DISETUJUI"]
    est_jasa = sum(i["subtotal"] for i in counted if i["source"] == "ESTIMASI" and i["kind"] == "jasa")
    est_part = sum(i["subtotal"] for i in counted if i["source"] == "ESTIMASI" and i["kind"] == "part")
    add_jasa = sum(i["subtotal"] for i in counted if i["source"] == "TAMBAHAN" and i["kind"] == "jasa")
    add_part = sum(i["subtotal"] for i in counted if i["source"] == "TAMBAHAN" and i["kind"] == "part")
    subtotal = est_jasa + est_part + add_jasa + add_part
    discount = int(trx.get("discount", 0))
    total = max(subtotal - discount, 0)
    # Modal part (dari cost_snapshot) & Laba per Nota
    modal_part = sum(int(i.get("cost", 0) or 0) * i["qty"] for i in counted if i["kind"] == "part")
    laba = max(total - modal_part, 0)
    trx["items"] = items
    trx["totals"] = {"est_jasa": est_jasa, "est_part": est_part, "add_jasa": add_jasa, "add_part": add_part,
                     "estimasi": est_jasa + est_part, "subtotal": subtotal, "discount": discount, "total": total,
                     "modal_part": modal_part, "laba": laba,
                     "pending_approval": sum(1 for i in items if i["approval"] == "MENUNGGU_PERSETUJUAN")}
    trx["customer"] = clean(await db.customers.find_one({"id": trx.get("customer_id")})) if trx.get("customer_id") else None
    trx["vehicle"] = clean(await db.vehicles.find_one({"id": trx.get("vehicle_id")})) if trx.get("vehicle_id") else None
    trx["payment"] = clean(await db.payments.find_one({"transaction_id": trx["id"]}))
    trx["invoice"] = clean(await db.invoices.find_one({"transaction_id": trx["id"]}))
    trx["status_logs"] = await db.transaction_status_logs.find({"transaction_id": trx["id"]}, {"_id": 0}).sort("created_at", 1).to_list(100)
    trx["audit_logs"] = await db.audit_logs.find({"transaction_id": trx["id"]}, {"_id": 0}).sort("created_at", 1).to_list(300)
    trx["whatsapp_logs"] = await db.whatsapp_logs.find({"transaction_id": trx["id"]}, {"_id": 0}).sort("created_at", -1).to_list(20)
    return trx


async def get_trx(tid: str) -> dict:
    trx = await db.service_transactions.find_one({"id": tid})
    if not trx:
        raise HTTPException(status_code=404, detail="Transaksi tidak ditemukan")
    return trx


async def ensure_editable(trx: dict, user: dict, owner_password: Optional[str] = None) -> bool:
    """Returns True when editing a locked (paid) transaction with owner authorization."""
    if trx["status"] == "DIBATALKAN":
        raise HTTPException(status_code=400, detail="Transaksi sudah dibatalkan")
    if trx["status"] in STATUS_PAID:
        if user["role"] == "owner" or await verify_owner_password(owner_password):
            return True
        raise HTTPException(status_code=400, detail="Transaksi sudah dikunci. Perlu konfirmasi password Owner untuk mengubah")
    return False


@api.post("/transactions")
async def create_transaction(body: TransactionCreate, user: RegistrarUser):
    customer = await db.customers.find_one({"id": body.customer_id})
    vehicle = await db.vehicles.find_one({"id": body.vehicle_id})
    if not customer or not vehicle:
        raise HTTPException(status_code=404, detail="Pelanggan atau motor tidak ditemukan")
    for it in body.items:
        if it.kind == "part" and it.ref_id:
            p = await db.parts.find_one({"id": it.ref_id, "deleted_at": None})
            if not p:
                raise HTTPException(status_code=404, detail="Part tidak ditemukan")
            if p["stock"] < it.qty:
                raise HTTPException(status_code=400, detail=f"STOK {'HABIS' if p['stock'] <= 0 else 'TIDAK CUKUP'}: {p['name']} (stok {p['stock']})")
            it.code = p["code"]
            if not it.name.strip():
                it.name = p["name"]
    today = now().astimezone(timezone(timedelta(hours=7))).strftime("%Y%m%d")
    created_at = iso(now())
    if body.date and user["role"] == "owner":
        bd = parse_date(body.date)
        if bd:
            today = bd.strftime("%Y%m%d")
            created_at = datetime(bd.year, bd.month, bd.day, 12, 0, tzinfo=timezone(timedelta(hours=7))).astimezone(timezone.utc).isoformat()
    q = await next_counter(f"queue:{today}")
    seq = await next_counter("trx")
    mech = await db.users.find_one({"id": body.mechanic_id}) if body.mechanic_id else None
    trx = {
        "id": new_id(), "trx_no": f"TRX-{today}-{seq:04d}", "queue_no": f"A-{q:03d}", "date": today,
        "customer_id": body.customer_id, "vehicle_id": body.vehicle_id, "customer_name": customer["name"], "customer_phone": customer.get("phone", ""),
        "plate": vehicle["plate"], "vehicle_name": f"{vehicle.get('brand', '')} {vehicle.get('model', '')}".strip(),
        "km_in": body.km_in, "complaint_main": body.complaint_main, "complaint_extra": body.complaint_extra, "initial_check": body.initial_check,
        "mechanic_id": mech["id"] if mech else None, "mechanic_name": mech["name"] if mech else None,
        "status": "MENUNGGU_SERVIS", "discount": 0, "work_done": "", "inspection": "", "mechanic_note": "", "condition": "",
        "next_recommendation": "", "next_km": None, "next_date": "", "started_at": None, "finished_at": None, "paid_at": None,
        "created_by": user["username"], "created_at": created_at, "updated_at": iso(now()), "deleted_at": None,
    }
    await db.service_transactions.insert_one(trx)
    await db.service_complaints.insert_one({"id": new_id(), "transaction_id": trx["id"], "main": body.complaint_main,
                                            "extra": body.complaint_extra, "initial_check": body.initial_check, "created_at": iso(now())})
    if body.items:
        docs = [item_doc(trx, it, "ESTIMASI", user, "DISETUJUI", await item_cost(it)) for it in body.items]
        await db.service_items.insert_many(docs)
        await db.service_estimations.insert_one({"id": new_id(), "transaction_id": trx["id"], "items": [clean(dict(d)) for d in docs],
                                                 "total": sum(d["subtotal"] for d in docs), "created_at": iso(now())})
    await log_status(trx["id"], "TERDAFTAR", user, "Pendaftaran servis")
    await log_status(trx["id"], "MENUNGGU_SERVIS", user, f"Nomor antrian {trx['queue_no']}")
    await audit(trx["id"], "REGISTER", user, f"{user['name']} mendaftarkan servis {trx['plate']} antrian {trx['queue_no']}")
    return await build_detail(trx)


@api.get("/transactions")
async def list_transactions(_: CurrentUser, status: str = "", date: str = "", q: str = "", mechanic_id: str = "", limit: int = 200):
    flt: dict[str, Any] = {"deleted_at": None}
    if status:
        groups = {"aktif": STATUS_ACTIVE, "kasir": ["MENUNGGU_KASIR", "MENUNGGU_PEMBAYARAN"], "dibayar": STATUS_PAID}
        flt["status"] = {"$in": groups.get(status, status.split(","))}
    if date:
        flt["date"] = date
    if mechanic_id:
        flt["mechanic_id"] = mechanic_id
    if q.strip():
        flt["$or"] = [{"trx_no": regex(q)}, {"queue_no": regex(q)}, {"plate": regex(q)}, {"customer_name": regex(q)},
                      {"customer_phone": regex(q)}, {"mechanic_name": regex(q)}, {"vehicle_name": regex(q)}]
    docs = await db.service_transactions.find(flt, {"_id": 0}).sort("created_at", -1).to_list(limit)
    ids = [d["id"] for d in docs]
    pipeline = [{"$match": {"transaction_id": {"$in": ids}, "deleted_at": None, "approval": "DISETUJUI"}},
                {"$group": {"_id": "$transaction_id", "subtotal": {"$sum": "$subtotal"}}}]
    sums = {r["_id"]: r["subtotal"] for r in await db.service_items.aggregate(pipeline).to_list(1000)}
    for d in docs:
        d["total"] = max(sums.get(d["id"], 0) - int(d.get("discount", 0)), 0)
    return docs


@api.get("/transactions/{tid}")
async def transaction_detail(tid: str, _: CurrentUser):
    return await build_detail(await get_trx(tid))


@api.post("/transactions/{tid}/start")
async def start_service(tid: str, body: StartIn, user: MekanikUser):
    trx = await get_trx(tid)
    if trx["status"] != "MENUNGGU_SERVIS":
        raise HTTPException(status_code=400, detail="Transaksi tidak dalam status MENUNGGU SERVIS")
    mech = await db.users.find_one({"id": body.mechanic_id}) if body.mechanic_id else None
    extra = {"started_at": iso(now()), "mechanic_id": mech["id"] if mech else user["id"], "mechanic_name": mech["name"] if mech else user["name"]}
    if body.km is not None:
        extra["km_in"] = body.km
    await set_status(tid, "DIPROSES", user, f"Mekanik {extra['mechanic_name']} mulai servis", extra)
    return await build_detail(await get_trx(tid))


@api.put("/transactions/{tid}/work")
async def update_work(tid: str, body: WorkNoteIn, user: MekanikUser):
    trx = await get_trx(tid)
    await ensure_editable(trx, user)
    await db.service_transactions.update_one({"id": tid}, {"$set": {**body.model_dump(), "updated_at": iso(now())}})
    await audit(tid, "WORK_NOTE", user, "Catatan pekerjaan diperbarui")
    return await build_detail(await get_trx(tid))


@api.post("/transactions/{tid}/items")
async def add_item(tid: str, body: ItemIn, user: RegistrarUser):
    trx = await get_trx(tid)
    locked = await ensure_editable(trx, user, body.owner_password)
    part = None
    if body.kind == "part":
        if not body.ref_id:
            raise HTTPException(status_code=400, detail="Pilih part dari master")
        part = await db.parts.find_one({"id": body.ref_id, "deleted_at": None})
        if not part:
            raise HTTPException(status_code=404, detail="Part tidak ditemukan")
        if part["stock"] <= 0:
            raise HTTPException(status_code=400, detail="STOK HABIS: part tidak dapat ditambahkan ke transaksi")
        active_ids = [t["id"] for t in await db.service_transactions.find({"status": {"$in": STATUS_ACTIVE}}, {"id": 1}).to_list(5000)]
        reserved = 0
        async for i in db.service_items.find({"ref_id": part["id"], "deleted_at": None, "approval": "DISETUJUI", "transaction_id": {"$in": active_ids}}):
            reserved += i["qty"]
        if part["stock"] - reserved < body.qty:
            raise HTTPException(status_code=400, detail=f"STOK TIDAK CUKUP: tersedia {max(part['stock'] - reserved, 0)} {part.get('unit', 'pcs')}")
        body.code, body.name = part["code"], part["name"]
    if locked:
        source, approval = "TAMBAHAN", "DISETUJUI"
    else:
        source = "ESTIMASI" if trx["status"] in ("TERDAFTAR", "MENUNGGU_SERVIS") else "TAMBAHAN"
        approval = "DISETUJUI" if source == "ESTIMASI" else "MENUNGGU_PERSETUJUAN"
    doc = item_doc(trx, body, source, user, approval, int(part.get("cost", 0) or 0) if part else 0)
    await db.service_items.insert_one(doc)
    if source == "TAMBAHAN":
        await db.service_additional_items.insert_one(dict(doc))
        if trx["status"] == "DIPROSES":
            await set_status(tid, "MENUNGGU_PERSETUJUAN", user, "Ada tambahan menunggu persetujuan")
    if locked:
        if part:
            await apply_stock_delta(part["id"], -body.qty, "KOREKSI FAKTUR", trx, user)
        await resync_paid(tid)
    await audit(tid, "ITEM_ADD" if not locked else "INVOICE_EDIT_ADD", user, f"Tambah {body.kind} {body.name} qty {body.qty} Rp{body.price * body.qty:,} ({source}){' [EDIT FAKTUR]' if locked else ''}")
    return await build_detail(await get_trx(tid))


@api.put("/transactions/{tid}/items/{iid}")
async def edit_item(tid: str, iid: str, body: ItemEditIn, user: RegistrarUser):
    trx = await get_trx(tid)
    locked = await ensure_editable(trx, user, body.owner_password)
    it = await db.service_items.find_one({"id": iid, "transaction_id": tid, "deleted_at": None})
    if not it:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    upd: dict[str, Any] = {}
    if body.name is not None and it["kind"] == "jasa" and body.name.strip():
        upd["name"] = body.name.strip()
    if body.price is not None and body.price >= 0:
        upd["price"] = body.price
    if body.qty is not None and body.qty > 0:
        upd["qty"] = body.qty
    if body.note is not None:
        upd["note"] = body.note
    new_price, new_qty = upd.get("price", it["price"]), upd.get("qty", it["qty"])
    if it["kind"] == "part" and it.get("ref_id") and new_qty != it["qty"]:
        p = await db.parts.find_one({"id": it["ref_id"]})
        extra_needed = new_qty - it["qty"]
        if p and extra_needed > 0 and p["stock"] < extra_needed and not locked:
            raise HTTPException(status_code=400, detail=f"STOK TIDAK CUKUP: tersedia {p['stock']}")
        if locked and p:
            if extra_needed > 0 and p["stock"] < extra_needed:
                raise HTTPException(status_code=400, detail=f"STOK TIDAK CUKUP: tersedia {p['stock']}")
            await apply_stock_delta(p["id"], -extra_needed, "KOREKSI FAKTUR", trx, user)
    upd["subtotal"] = new_price * new_qty
    await db.service_items.update_one({"id": iid}, {"$set": upd})
    changes = ", ".join(f"{k}: {it.get(k)} → {v}" for k, v in upd.items() if k != "subtotal" and it.get(k) != v)
    await audit(tid, "INVOICE_EDIT" if locked else "ITEM_EDIT", user, f"Ubah item {it['name']} ({changes}){'. Alasan: ' + body.reason if body.reason else ''}{' [EDIT FAKTUR]' if locked else ''}",
                {"before": {k: it.get(k) for k in ("name", "price", "qty")}, "after": {k: upd.get(k, it.get(k)) for k in ("name", "price", "qty")}, "reason": body.reason})
    if locked:
        await resync_paid(tid)
    return await build_detail(await get_trx(tid))


async def maybe_resume(tid: str, user: dict):
    pending = await db.service_items.count_documents({"transaction_id": tid, "deleted_at": None, "approval": "MENUNGGU_PERSETUJUAN"})
    trx = await get_trx(tid)
    if pending == 0 and trx["status"] == "MENUNGGU_PERSETUJUAN":
        await set_status(tid, "DIPROSES", user, "Semua tambahan sudah ditinjau")


@api.post("/transactions/{tid}/items/{iid}/approve")
async def approve_item(tid: str, iid: str, user: RegistrarUser):
    trx = await get_trx(tid)
    await ensure_editable(trx, user)
    it = await db.service_items.find_one_and_update({"id": iid, "transaction_id": tid}, {"$set": {"approval": "DISETUJUI", "approved_by": user["username"]}}, return_document=True)
    if not it:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await audit(tid, "ITEM_APPROVE", user, f"Setujui tambahan {it['name']}")
    await maybe_resume(tid, user)
    return await build_detail(await get_trx(tid))


@api.post("/transactions/{tid}/items/{iid}/reject")
async def reject_item(tid: str, iid: str, user: RegistrarUser):
    trx = await get_trx(tid)
    await ensure_editable(trx, user)
    it = await db.service_items.find_one_and_update({"id": iid, "transaction_id": tid}, {"$set": {"approval": "DITOLAK", "approved_by": user["username"]}}, return_document=True)
    if not it:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await audit(tid, "ITEM_REJECT", user, f"Tolak tambahan {it['name']}")
    await maybe_resume(tid, user)
    return await build_detail(await get_trx(tid))


@api.delete("/transactions/{tid}/items/{iid}")
async def delete_item(tid: str, iid: str, user: RegistrarUser, owner_password: str = ""):
    trx = await get_trx(tid)
    locked = await ensure_editable(trx, user, owner_password or None)
    it = await db.service_items.find_one_and_update({"id": iid, "transaction_id": tid, "deleted_at": None}, {"$set": {"deleted_at": iso(now())}}, return_document=True)
    if not it:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    if locked:
        if it["kind"] == "part" and it.get("ref_id") and it["approval"] == "DISETUJUI":
            await apply_stock_delta(it["ref_id"], it["qty"], "KOREKSI FAKTUR", trx, user)
        await resync_paid(tid)
    await audit(tid, "INVOICE_EDIT_REMOVE" if locked else "ITEM_REMOVE", user, f"Hapus item {it['name']}{' [EDIT FAKTUR]' if locked else ''}")
    await maybe_resume(tid, user)
    return await build_detail(await get_trx(tid))


@api.put("/transactions/{tid}/items/{iid}/price")
async def change_price(tid: str, iid: str, body: PriceChangeIn, user: OwnerUser):
    it = await db.service_items.find_one({"id": iid, "transaction_id": tid})
    if not it:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    await db.service_items.update_one({"id": iid}, {"$set": {"price": body.price, "subtotal": body.price * it["qty"]}})
    await audit(tid, "PRICE_CHANGE", user, f"Harga {it['name']} Rp{it['price']:,} → Rp{body.price:,}. Alasan: {body.reason}",
                {"before": it["price"], "after": body.price, "reason": body.reason})
    return await build_detail(await get_trx(tid))


@api.post("/transactions/{tid}/finish")
async def finish_service(tid: str, body: FinishIn, user: MekanikUser):
    trx = await get_trx(tid)
    if trx["status"] not in ("DIPROSES", "MENUNGGU_PERSETUJUAN"):
        raise HTTPException(status_code=400, detail="Servis belum dimulai atau sudah selesai")
    pending = await db.service_items.count_documents({"transaction_id": tid, "deleted_at": None, "approval": "MENUNGGU_PERSETUJUAN"})
    if pending:
        raise HTTPException(status_code=400, detail=f"Masih ada {pending} tambahan menunggu persetujuan")
    # SOP Mekanik: checklist wajib 100% OK sebelum servis bisa diselesaikan
    req = await sop_required_names()
    if req:
        done = {c["name"]: c.get("ok", False) for c in (trx.get("sop_checklist") or [])}
        missing = [n for n in req if not done.get(n)]
        if missing:
            raise HTTPException(status_code=400, detail="Checklist SOP belum lengkap: " + ", ".join(missing))
    extra = {**body.model_dump(), "finished_at": iso(now())}
    await set_status(tid, "MENUNGGU_KASIR", user, "Pekerjaan selesai, Work Order diserahkan ke kasir", extra)
    detail = await build_detail(await get_trx(tid))
    await db.work_orders.update_one({"transaction_id": tid}, {"$set": {"id": new_id(), "transaction_id": tid, "trx_no": detail["trx_no"],
                                    "queue_no": detail["queue_no"], "snapshot": {k: v for k, v in detail.items() if k in ("customer", "vehicle", "items", "totals", "mechanic_name", "complaint_main")},
                                    "created_at": iso(now())}}, upsert=True)
    await audit(tid, "WORK_ORDER", user, "Work Order dikirim ke kasir")
    return detail


@api.post("/transactions/{tid}/checkout")
async def checkout(tid: str, user: KasirUser):
    trx = await get_trx(tid)
    if trx["status"] == "MENUNGGU_KASIR":
        await set_status(tid, "MENUNGGU_PEMBAYARAN", user, "Kasir menerima Work Order")
    return await build_detail(await get_trx(tid))


@api.post("/transactions/{tid}/pay")
async def pay(tid: str, body: PaymentIn, user: KasirUser):
    # Atomic status lock prevents double-submit and double stock deduction.
    trx = await db.service_transactions.find_one_and_update(
        {"id": tid, "status": {"$in": ["MENUNGGU_KASIR", "MENUNGGU_PEMBAYARAN"]}},
        {"$set": {"status": "PEMBAYARAN_DIPROSES", "discount": body.discount, "updated_at": iso(now())}}, return_document=True)
    if not trx:
        raise HTTPException(status_code=400, detail="Transaksi tidak dalam status menunggu pembayaran")
    detail = await build_detail(dict(trx))
    total = detail["totals"]["total"]
    if body.method != "KREDIT" and body.amount_paid < total:
        await db.service_transactions.update_one({"id": tid}, {"$set": {"status": "MENUNGGU_PEMBAYARAN"}})
        raise HTTPException(status_code=400, detail="Jumlah dibayar kurang dari total tagihan")
    part_items = [i for i in detail["items"] if i["kind"] == "part" and i["approval"] == "DISETUJUI" and i.get("ref_id")]
    for i in part_items:
        p = await db.parts.find_one({"id": i["ref_id"]})
        if p and p["stock"] < i["qty"]:
            await db.service_transactions.update_one({"id": tid}, {"$set": {"status": "MENUNGGU_PEMBAYARAN"}})
            raise HTTPException(status_code=400, detail=f"STOK TIDAK CUKUP untuk {p['name']} (stok {p['stock']})")
    for i in part_items:
        await apply_stock_delta(i["ref_id"], -i["qty"], "SERVIS", trx, user)
    change = max(body.amount_paid - total, 0) if body.method == "CASH" else 0
    debt = max(total - body.amount_paid, 0) if body.method == "KREDIT" else 0
    due_norm = ""
    if debt:
        dd = parse_date(body.due_date)
        due_norm = dd.isoformat() if dd else (now().astimezone(timezone(timedelta(hours=7))).date() + timedelta(days=7)).isoformat()
    payment = {"id": new_id(), "transaction_id": tid, "method": body.method, "total": total, "amount_paid": body.amount_paid, "change": change,
               "discount": body.discount, "reference": body.reference, "note": body.note, "cashier": user["name"], "paid_at": iso(now()), "date": wib_today(),
               "due_date": due_norm, "installments": []}
    await db.payments.insert_one(payment)
    seq = await next_counter("invoice")
    inv_no = f"NS-{trx['date']}-{seq:04d}"
    await db.invoices.insert_one({"id": new_id(), "transaction_id": tid, "invoice_no": inv_no, "total": total, "printed_at": None, "created_at": iso(now())})
    await set_status(tid, "SUDAH_DIBAYAR", user, f"Pembayaran {body.method} Rp{body.amount_paid:,} dikonfirmasi" + (f" (hutang Rp{debt:,})" if debt else ""),
                     {"paid_at": iso(now()), "invoice_no": inv_no, "locked": True, "debt_amount": debt, "debt_due_date": due_norm,
                      "debt_status": "BELUM LUNAS" if debt else ("LUNAS" if body.method == "KREDIT" else None)})
    if debt:
        kind = "HUTANG PENUH" if body.amount_paid == 0 else "HUTANG SEBAGIAN"
        await notify_owner("HUTANG", f"{kind}: {trx['customer_name']}",
                           f"Nota {inv_no} · {trx['plate']}\nNominal hutang: Rp {debt:,}".replace(",", ".") + f"\nJatuh tempo: {body.due_date or '-'}\nKasir: {user['name']}",
                           tid, {"customer_name": trx["customer_name"], "amount": debt, "due_date": body.due_date, "invoice_no": inv_no})
    await db.vehicles.update_one({"id": trx["vehicle_id"]}, {"$set": {"last_service_at": iso(now()), "km_last": trx.get("km_in", 0)}})
    await db.service_history.update_one({"transaction_id": tid}, {"$set": {"id": new_id(), "transaction_id": tid, "vehicle_id": trx["vehicle_id"],
                                        "customer_id": trx["customer_id"], "plate": trx["plate"], "invoice_no": inv_no, "total": total,
                                        "mechanic_name": trx.get("mechanic_name"), "date": trx["date"], "created_at": iso(now())}}, upsert=True)
    await audit(tid, "PAYMENT", user, f"Kasir {user['name']} konfirmasi pembayaran {body.method} Rp{total:,}")
    return await build_detail(await get_trx(tid))


@api.post("/transactions/{tid}/pay-debt")
async def pay_debt(tid: str, body: DebtPayIn, user: KasirUser):
    trx = await get_trx(tid)
    debt = int(trx.get("debt_amount") or 0)
    if trx["status"] not in STATUS_PAID or debt <= 0:
        raise HTTPException(status_code=400, detail="Transaksi tidak memiliki hutang")
    if body.amount <= 0:
        raise HTTPException(status_code=400, detail="Nominal pembayaran tidak valid")
    amount = min(body.amount, debt)
    new_debt = debt - amount
    inst = {"id": new_id(), "amount": amount, "method": body.method, "reference": body.reference, "note": body.note, "cashier": user["name"], "paid_at": iso(now()), "date": wib_today()}
    await db.payments.update_one({"transaction_id": tid}, {"$inc": {"amount_paid": amount}, "$push": {"installments": inst}})
    await db.service_transactions.update_one({"id": tid}, {"$set": {"debt_amount": new_debt, "debt_status": "LUNAS" if new_debt == 0 else "BELUM LUNAS", "updated_at": iso(now())}})
    await db.debt_payments.insert_one({**inst, "transaction_id": tid, "debt_before": debt, "debt_after": new_debt})
    await audit(tid, "DEBT_PAYMENT", user, f"Pembayaran hutang {body.method} Rp{amount:,}. Sisa hutang Rp{new_debt:,}")
    if new_debt == 0:
        await notify_owner("HUTANG LUNAS", f"HUTANG LUNAS: {trx['customer_name']}", f"Nota {trx.get('invoice_no')} · {trx['plate']}\nHutang telah dilunasi (Rp {amount:,}).".replace(",", "."), tid)
    return await build_detail(await get_trx(tid))


@api.get("/debts")
async def list_debts(_: KasirUser, status: str = "BELUM LUNAS"):
    today = now().astimezone(timezone(timedelta(hours=7))).date()
    flt: dict[str, Any] = {"debt_status": {"$ne": None}, "deleted_at": None}
    if status:
        flt["debt_status"] = status
    docs = await db.service_transactions.find(flt, {"_id": 0}).sort("paid_at", -1).to_list(500)
    for d in docs:
        due = parse_date(d.get("debt_due_date") or "")
        d["overdue"] = bool(due and due < today and d.get("debt_status") == "BELUM LUNAS")
        d["sale"] = False
    sales = await db.sales.find(flt, {"_id": 0}).sort("created_at", -1).to_list(500)
    for s in sales:
        due = parse_date(s.get("debt_due_date") or "")
        s["overdue"] = bool(due and due < today and s.get("debt_status") == "BELUM LUNAS")
        s["sale"] = True
        s["invoice_no"] = s.get("sale_no")
        s["plate"] = "PENJUALAN PART"
        s["paid_at"] = s.get("created_at")
    merged = docs + sales
    merged.sort(key=lambda x: (x.get("paid_at") or x.get("created_at") or ""), reverse=True)
    return merged


@api.post("/transactions/{tid}/print")
async def mark_printed(tid: str, user: KasirUser):
    trx = await get_trx(tid)
    if trx["status"] not in STATUS_PAID:
        raise HTTPException(status_code=400, detail="Transaksi belum dibayar")
    req = await sop_required_names()
    if req:
        done = {c["name"]: c.get("ok", False) for c in (trx.get("sop_checklist") or [])}
        missing = [n for n in req if not done.get(n)]
        if missing:
            raise HTTPException(status_code=400, detail="Checklist SOP belum lengkap, nota tidak dapat dicetak: " + ", ".join(missing))
    await db.invoices.update_one({"transaction_id": tid}, {"$set": {"printed_at": iso(now())}})
    if trx["status"] == "SUDAH_DIBAYAR":
        await set_status(tid, "NOTA_DICETAK", user, "Nota dicetak")
    else:
        await audit(tid, "PRINT", user, "Nota dicetak ulang")
    return await build_detail(await get_trx(tid))


def wa_message(d: dict, shop: dict) -> str:
    inv = d.get("invoice") or {}
    dt = (d.get("paid_at") or d["created_at"])[:10]
    y, m, dd = dt.split("-")
    name = shop["name"]
    return (f"Terima kasih telah melakukan servis di {name}.\n\n"
            f"No Nota: {inv.get('invoice_no', d['trx_no'])}\nTanggal: {dd}/{m}/{y}\nPelanggan: {d['customer_name']}\n"
            f"Motor: {d.get('vehicle_name') or '-'}\nNo Polisi: {d['plate']}\nTotal: Rp {d['totals']['total']:,}".replace(",", ".") +
            (f"\nSisa hutang: Rp {d.get('debt_amount', 0):,}".replace(",", ".") + (f" (jatuh tempo {d.get('debt_due_date')})" if d.get("debt_due_date") else "") if d.get("debt_amount") else "") +
            f"\n\nTerima kasih telah mempercayakan perawatan motor Anda kepada {name}.")


@api.post("/transactions/{tid}/whatsapp")
async def send_whatsapp(tid: str, user: KasirUser):
    trx = await get_trx(tid)
    if trx["status"] not in STATUS_PAID:
        raise HTTPException(status_code=400, detail="Transaksi belum dibayar")
    d = await build_detail(dict(trx))
    phone = wa_phone(d.get("customer_phone") or "")
    msg = wa_message(d, await get_shop())
    status = "TERKIRIM" if phone else "GAGAL"
    await db.whatsapp_logs.insert_one({"id": new_id(), "transaction_id": tid, "phone": phone, "message": msg, "status": status,
                                       "sent_by": user["username"], "created_at": iso(now())})
    if status == "TERKIRIM" and trx["status"] in ("SUDAH_DIBAYAR", "NOTA_DICETAK"):
        await set_status(tid, "NOTA_TERKIRIM", user, "Nota dikirim via WhatsApp")
    else:
        await audit(tid, "WHATSAPP", user, f"Kirim WhatsApp: {status}")
    from urllib.parse import quote
    return {"status": status, "phone": phone, "message": msg, "url": f"https://wa.me/{phone}?text={quote(msg)}" if phone else None}


@api.post("/transactions/{tid}/complete")
async def complete(tid: str, user: KasirUser):
    trx = await get_trx(tid)
    if trx["status"] not in STATUS_PAID:
        raise HTTPException(status_code=400, detail="Transaksi belum dibayar")
    if trx["status"] != "SELESAI":
        await set_status(tid, "SELESAI", user, "Transaksi selesai")
    return await build_detail(await get_trx(tid))


@api.post("/transactions/{tid}/cancel")
async def cancel(tid: str, body: CancelIn, user: RegistrarUser):
    trx = await get_trx(tid)
    if trx["status"] == "DIBATALKAN":
        raise HTTPException(status_code=400, detail="Transaksi sudah dibatalkan")
    # KEAMANAN: SEMUA pembatalan wajib konfirmasi password Owner (termasuk saat login sebagai Owner).
    if not await verify_owner_password(body.owner_password):
        raise HTTPException(status_code=403, detail="Pembatalan wajib konfirmasi password Owner")
    reason = (body.reason or "").strip()
    if not reason:
        raise HTTPException(status_code=400, detail="Alasan pembatalan wajib diisi")
    was_paid = trx["status"] in STATUS_PAID
    if was_paid:
        # restore stock of used parts
        items = await db.service_items.find({"transaction_id": tid, "deleted_at": None, "approval": "DISETUJUI", "kind": "part"}).to_list(500)
        for i in items:
            if i.get("ref_id"):
                await apply_stock_delta(i["ref_id"], i["qty"], "PEMBATALAN", trx, user)
    # Pembayaran ditandai cancelled agar tidak mencemari Omset/Laba/Dashboard
    await db.payments.update_many({"transaction_id": tid}, {"$set": {"cancelled": True}})
    await db.service_history.update_one({"transaction_id": tid}, {"$set": {"cancelled": True}})
    await set_status(tid, "DIBATALKAN", user, f"Dibatalkan: {reason}",
                     {"cancel_reason": reason, "cancelled_by": user["name"], "cancelled_at": iso(now()), "was_paid": was_paid, "status_before_cancel": trx["status"]})
    return await build_detail(await get_trx(tid))


@api.get("/transactions/{tid}/receipt")
async def receipt(tid: str, _: CurrentUser):
    d = await build_detail(await get_trx(tid))
    shop = await get_shop()
    return {"shop": shop, "detail": d, "whatsapp_message": wa_message(d, shop)}


# --------------------------------------------------------------------------- history

@api.get("/history")
async def history(_: CurrentUser, q: str = "", vehicle_id: str = ""):
    flt: dict[str, Any] = {"status": {"$in": STATUS_PAID}, "deleted_at": None}
    if vehicle_id:
        flt["vehicle_id"] = vehicle_id
    if q.strip():
        flt["$or"] = [{"trx_no": regex(q)}, {"invoice_no": regex(q)}, {"plate": regex(q)}, {"customer_name": regex(q)},
                      {"customer_phone": regex(q)}, {"mechanic_name": regex(q)}, {"vehicle_name": regex(q)}, {"date": regex(q)}]
    docs = await db.service_transactions.find(flt, {"_id": 0}).sort("created_at", -1).to_list(300)
    ids = [d["id"] for d in docs]
    items = await db.service_items.find({"transaction_id": {"$in": ids}, "deleted_at": None, "approval": "DISETUJUI"}, {"_id": 0}).to_list(3000)
    by = {}
    for i in items:
        by.setdefault(i["transaction_id"], []).append(i)
    for d in docs:
        its = by.get(d["id"], [])
        d["item_names"] = [i["name"] for i in its]
        d["total"] = max(sum(i["subtotal"] for i in its) - int(d.get("discount", 0)), 0)
    return docs


# --------------------------------------------------------------------------- dashboard & reports


@api.get("/reports/dashboard-bengkel")
async def dashboard_bengkel(user: OwnerUser):
    """Dashboard Bengkel baru: Omzet (hanya servis status SELESAI/SUDAH_DIBAYAR, bukan DIBATALKAN),
    Modal Part, Belanja Bengkel, Laba Bersih hari ini + tren 7 hari. Prive Keluarga dipisah, tidak masuk laba."""
    wib_now = now().astimezone(WIB).date()
    day_keys = [(wib_now - timedelta(days=i)).strftime("%Y%m%d") for i in range(6, -1, -1)]  # 7 hari, ASC

    pays = await db.payments.find({"cancelled": {"$ne": True}, "sale": {"$ne": True}, "date": {"$in": day_keys}}, {"_id": 0}).to_list(50000)
    modal_map = await _modal_by_payment(pays)
    sales = await db.sales.find({"deleted_at": None, "date": {"$in": day_keys}}, {"_id": 0}).to_list(50000)
    exp = await db.expenses.find({"deleted_at": None, "date": {"$in": day_keys}}, {"_id": 0}).to_list(50000)

    trend: dict[str, dict] = {d: {"omzet": 0, "modal": 0, "belanja_bengkel": 0, "prive": 0, "laba_bersih": 0} for d in day_keys}
    for p in pays:
        b = trend.get(str(p.get("date")))
        if b:
            b["omzet"] += int(p.get("total", 0) or 0)
            b["modal"] += modal_map.get(p["transaction_id"], 0)
    for s in sales:
        b = trend.get(str(s.get("date")))
        if b:
            b["omzet"] += int(s.get("total", 0) or 0)
            b["modal"] += int(s.get("total_cost", 0) or 0)
    for e in exp:
        b = trend.get(str(e.get("date")))
        if not b:
            continue
        amt = int(e.get("amount", 0) or 0)
        if e.get("group") == "BENGKEL":
            b["belanja_bengkel"] += amt if e.get("flow", "OUT") == "OUT" else -amt
        elif e.get("group") == "KELUARGA":
            b["prive"] += amt
    rows = []
    for d in day_keys:
        b = trend[d]
        b["laba_bersih"] = b["omzet"] - b["modal"] - b["belanja_bengkel"]  # Prive TIDAK masuk laba
        rows.append({"date": f"{d[:4]}-{d[4:6]}-{d[6:8]}", **b})
    t = trend[wib_today()]
    # Ringkasan bulanan + Target Laba + Pengingat Kas Tipis
    month = wib_now.strftime("%Y-%m")
    lr_month = await build_laba_rugi(month)
    laba_month = int(lr_month["total"]["laba_bersih"])
    fin = await db.settings.find_one({"id": "finance"}, {"_id": 0}) or {}
    target = int(fin.get("target_laba_bulanan", 0) or 0)
    target_pct = round((laba_month / target) * 100, 1) if target > 0 else None
    cf = await report_cashflow(user, month)
    saldo_bengkel = int(cf["bengkel"]["saldo_akhir"])
    belanja_month = int(cf["bengkel"]["pengeluaran"])
    if saldo_bengkel <= 0:
        kas_level, kas_msg = "danger", "Kas bengkel habis / minus. Segera setor modal atau tahan pengeluaran."
    elif belanja_month >= saldo_bengkel:
        kas_level, kas_msg = "danger", "Belanja bengkel sudah melebihi sisa kas. Hentikan pengeluaran tidak penting."
    elif belanja_month >= saldo_bengkel * 0.8:
        kas_level, kas_msg = "warning", "Belanja bengkel hampir melebihi saldo kas. Hati-hati atur pengeluaran."
    else:
        kas_level, kas_msg = "ok", "Kas bengkel masih aman."
    return {
        "today": {"date": wib_now.isoformat(), "omzet": t["omzet"], "modal_part": t["modal"],
                  "belanja_bengkel": t["belanja_bengkel"], "laba_bersih": t["laba_bersih"], "prive": t["prive"]},
        "trend": rows,
        "prive_7d": sum(r["prive"] for r in rows),
        "month": {"period": month, "laba_bersih": laba_month, "target": target, "target_pct": target_pct,
                  "tercapai": target > 0 and laba_month >= target},
        "kas": {"saldo_bengkel": saldo_bengkel, "belanja_bulan": belanja_month, "level": kas_level, "message": kas_msg},
    }


@api.get("/dashboard")
async def dashboard(user: CurrentUser):
    today = wib_today()
    trx_today = await db.service_transactions.find({"date": today, "deleted_at": None}, {"_id": 0}).to_list(1000)
    count = lambda *st: sum(1 for t in trx_today if t["status"] in st)  # noqa: E731
    pays = await db.payments.find({"date": today, "cancelled": {"$ne": True}}, {"_id": 0}).to_list(1000)
    paid_ids = [p["transaction_id"] for p in pays]
    items = await db.service_items.find({"transaction_id": {"$in": paid_ids}, "deleted_at": None, "approval": "DISETUJUI"}, {"_id": 0}).to_list(5000)
    by_method: dict[str, int] = {}
    for p in pays:
        by_method[p["method"]] = by_method.get(p["method"], 0) + p["total"]
    month_prefix = today[:6]
    month_pays = await db.payments.find({"date": {"$regex": f"^{month_prefix}"}, "cancelled": {"$ne": True}}, {"_id": 0, "total": 1}).to_list(10000)
    low_stock = await db.parts.count_documents({"deleted_at": None, "$expr": {"$lte": ["$stock", "$min_stock"]}})
    unread = await unread_count(user)
    tools_status = await tools_due_status(user) if user["role"] in ("mekanik", "owner") else None
    debts = await db.service_transactions.find({"debt_status": "BELUM LUNAS"}, {"_id": 0, "debt_amount": 1}).to_list(1000)
    reminders_due = len([r for r in await compute_reminders() if r["status"] in ("TERLAMBAT", "SEGERA")])
    debt_reminders_due = len(await debt_reminders(user))
    last_backup = await db.meta.find_one({"_id": "last_backup"}) or {}
    perf_pipe = [{"$match": {"date": today, "status": {"$in": ["MENUNGGU_KASIR", "MENUNGGU_PEMBAYARAN", *STATUS_PAID]}}},
                 {"$group": {"_id": "$mechanic_name", "count": {"$sum": 1}}}]
    perf = [{"mechanic": r["_id"] or "-", "count": r["count"]} for r in await db.service_transactions.aggregate(perf_pipe).to_list(50)]
    return {
        "today": today, "role": user["role"],
        "queue_today": count(*STATUS_ACTIVE), "menunggu": count("MENUNGGU_SERVIS"), "diproses": count("DIPROSES"),
        "menunggu_persetujuan": count("MENUNGGU_PERSETUJUAN"), "menunggu_kasir": count("MENUNGGU_KASIR"),
        "menunggu_pembayaran": count("MENUNGGU_PEMBAYARAN"), "sudah_dibayar": count(*STATUS_PAID), "selesai": count("SELESAI"),
        "total_servis_hari_ini": len(trx_today), "transaksi_hari_ini": len(pays),
        "omzet_hari_ini": sum(p["total"] for p in pays), "omzet_bulan_ini": sum(p["total"] for p in month_pays),
        "omzet_jasa": sum(i["subtotal"] for i in items if i["kind"] == "jasa"), "omzet_part": sum(i["subtotal"] for i in items if i["kind"] == "part"),
        "by_method": by_method, "hutang": by_method.get("KREDIT", 0),
        "jumlah_pelanggan": await db.customers.count_documents({"deleted_at": None}),
        "jumlah_kendaraan": await db.vehicles.count_documents({"deleted_at": None}), "stok_menipis": low_stock, "performa_mekanik": perf,
        "notif_unread": unread, "piutang_count": len(debts), "piutang_total": sum(int(d.get("debt_amount") or 0) for d in debts), "reminders_due": reminders_due, "debt_reminders_due": debt_reminders_due,
        "tools": tools_status,
        "backup_today": last_backup.get("date") == today, "last_backup_at": last_backup.get("at"),
    }


@api.get("/reports/omzet")
async def report_omzet(_: OwnerUser, mode: Literal["daily", "monthly"] = "daily"):
    pays = await db.payments.find({"cancelled": {"$ne": True}}, {"_id": 0}).to_list(50000)
    ids = [p["transaction_id"] for p in pays]
    items = await db.service_items.find({"transaction_id": {"$in": ids}, "deleted_at": None, "approval": "DISETUJUI"}, {"_id": 0}).to_list(100000)
    jasa: dict[str, int] = {}
    part: dict[str, int] = {}
    for i in items:
        (jasa if i["kind"] == "jasa" else part)[i["transaction_id"]] = (jasa if i["kind"] == "jasa" else part).get(i["transaction_id"], 0) + i["subtotal"]
    buckets: dict[str, dict] = {}
    for p in pays:
        local = datetime.fromisoformat(p["paid_at"]).astimezone(timezone(timedelta(hours=7)))
        key = local.strftime("%Y-%m-%d") if mode == "daily" else local.strftime("%Y-%m")
        b = buckets.setdefault(key, {"period": key, "total": 0, "jasa": 0, "part": 0, "count": 0, "diskon": 0})
        b["total"] += p["total"]
        b["jasa"] += jasa.get(p["transaction_id"], 0)
        b["part"] += part.get(p["transaction_id"], 0)
        b["diskon"] += p.get("discount", 0)
        b["count"] += 1
    rows = sorted(buckets.values(), key=lambda r: r["period"], reverse=True)[: (30 if mode == "daily" else 12)]
    return {"mode": mode, "rows": rows, "grand_total": sum(r["total"] for r in rows)}


@api.get("/reports/profit-trend")
async def profit_trend(_: OwnerUser):
    wib = timezone(timedelta(hours=7))
    today = now().astimezone(wib).date()
    # 12 bulan terakhir (key YYYY-MM), urut lama -> baru
    keys: list[str] = []
    y, m = today.year, today.month
    for _ in range(12):
        keys.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            m, y = 12, y - 1
    keys = list(reversed(keys))

    # laba part dari servis: butuh cost per item (fallback by code / ref_id)
    items = await db.service_items.find({"kind": "part", "deleted_at": None, "approval": "DISETUJUI"}, {"_id": 0}).to_list(200000)
    codes = list({i.get("code") for i in items if i.get("code")})
    ref_ids = list({i.get("ref_id") for i in items if i.get("ref_id")})
    cost_by_code = {p["code"]: int(p.get("cost", 0) or 0) for p in await db.parts.find({"code": {"$in": codes}}, {"_id": 0, "code": 1, "cost": 1}).to_list(200000)}
    cost_by_id = {p["id"]: int(p.get("cost", 0) or 0) for p in await db.parts.find({"id": {"$in": ref_ids}}, {"_id": 0, "id": 1, "cost": 1}).to_list(200000)}
    svc_part_profit: dict[str, int] = {}
    svc_part_omzet: dict[str, int] = {}
    for i in items:
        cost = int(i.get("cost") or 0) or cost_by_id.get(i.get("ref_id"), 0) or cost_by_code.get(i.get("code"), 0)
        svc_part_profit[i["transaction_id"]] = svc_part_profit.get(i["transaction_id"], 0) + (i["subtotal"] - cost * i["qty"])
        svc_part_omzet[i["transaction_id"]] = svc_part_omzet.get(i["transaction_id"], 0) + i["subtotal"]

    monthly: dict[str, dict] = {}

    def bucket(k: str) -> dict:
        return monthly.setdefault(k, {"period": k, "part_profit": 0, "service_part_profit": 0, "sales_profit": 0, "part_omzet": 0})

    # servis: pakai tanggal bayar (paid_at)
    for p in await db.payments.find({"sale": {"$ne": True}, "cancelled": {"$ne": True}}, {"_id": 0, "transaction_id": 1, "paid_at": 1}).to_list(100000):
        tid = p.get("transaction_id")
        if tid not in svc_part_profit and tid not in svc_part_omzet:
            continue
        try:
            k = datetime.fromisoformat(p["paid_at"]).astimezone(wib).strftime("%Y-%m")
        except (ValueError, TypeError, KeyError):
            continue
        b = bucket(k)
        b["service_part_profit"] += svc_part_profit.get(tid, 0)
        b["part_profit"] += svc_part_profit.get(tid, 0)
        b["part_omzet"] += svc_part_omzet.get(tid, 0)

    # jualan langsung
    for s in await db.sales.find({"deleted_at": None}, {"_id": 0, "date": 1, "total_profit": 1, "total": 1}).to_list(100000):
        d = str(s.get("date", ""))
        if len(d) < 6:
            continue
        k = f"{d[:4]}-{d[4:6]}"
        b = bucket(k)
        b["sales_profit"] += int(s.get("total_profit", 0) or 0)
        b["part_profit"] += int(s.get("total_profit", 0) or 0)
        b["part_omzet"] += int(s.get("total", 0) or 0)

    rows = [monthly.get(k, {"period": k, "part_profit": 0, "service_part_profit": 0, "sales_profit": 0, "part_omzet": 0}) for k in keys]

    # tahunan
    yearly: dict[str, int] = {}
    for b in monthly.values():
        yr = b["period"][:4]
        yearly[yr] = yearly.get(yr, 0) + b["part_profit"]
    yearly_rows = [{"year": yr, "part_profit": yearly[yr]} for yr in sorted(yearly.keys(), reverse=True)]

    # perbandingan bulan ini vs bulan lalu (M-1)
    cur_k = keys[-1]
    prev_k = keys[-2]
    cur = monthly.get(cur_k, {}).get("part_profit", 0)
    prev = monthly.get(prev_k, {}).get("part_profit", 0)
    delta = cur - prev
    pct = round((delta / prev) * 100, 1) if prev else (100.0 if cur else 0.0)
    this_year = f"{today.year:04d}"
    return {
        "rows": rows,
        "yearly": yearly_rows,
        "this_year": this_year,
        "this_year_profit": yearly.get(this_year, 0),
        "mom": {"current_period": cur_k, "prev_period": prev_k, "current": cur, "prev": prev, "delta": delta, "pct": pct},
    }


# --------------------------------------------------------------------------- excel export / import

EXPORT_SPECS = {
    "customers": {"coll": "customers", "sheet": "Pelanggan", "cols": [("code", "Kode"), ("name", "Nama"), ("phone", "No HP"), ("address", "Alamat"), ("notes", "Catatan")]},
    "services": {"coll": "services", "sheet": "Jasa", "cols": [("code", "Kode"), ("name", "Nama Jasa"), ("price", "Harga"), ("active", "Aktif")]},
    "parts": {"coll": "parts", "sheet": "Stok Barang", "cols": [("code", "Kode"), ("name", "Nama Part"), ("barcode", "Barcode"), ("price", "Harga Jual"), ("cost", "Harga Beli"), ("stock", "Stok"), ("min_stock", "Stok Minimum"), ("unit", "Satuan"), ("rack", "Lokasi Rak")]},
    "vehicles": {"coll": "vehicles", "sheet": "Motor", "cols": [("plate", "No Polisi"), ("customer_phone", "No HP Pelanggan"), ("customer_name", "Nama Pelanggan"), ("brand", "Merek"), ("model", "Tipe"), ("year", "Tahun"), ("color", "Warna"), ("chassis_no", "No Rangka"), ("engine_no", "No Mesin"), ("km_last", "KM Terakhir")]},
}


@api.get("/export/{entity}")
async def export_excel(entity: str, _: OwnerUser):
    spec = EXPORT_SPECS.get(entity)
    if not spec:
        raise HTTPException(status_code=404, detail="Jenis data tidak dikenal")
    docs = await db[spec["coll"]].find({"deleted_at": None}, {"_id": 0}).to_list(50000)
    if entity == "vehicles":
        custs = {c["id"]: c for c in await db.customers.find({}, {"_id": 0}).to_list(50000)}
        for d in docs:
            c = custs.get(d["customer_id"], {})
            d["customer_phone"], d["customer_name"] = c.get("phone", ""), c.get("name", "")
    wb = Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    ws.append([h for _, h in spec["cols"]])
    for d in docs:
        ws.append([d.get(k, "") for k, _ in spec["cols"]])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"suel_{entity}_{wib_today()}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@api.post("/import/{entity}")
async def import_excel(entity: str, user: OwnerUser, file: UploadFile = File(...)):
    spec = EXPORT_SPECS.get(entity)
    if not spec:
        raise HTTPException(status_code=404, detail="Jenis data tidak dikenal")
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="File bukan Excel (.xlsx) yang valid")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="File kosong")
    header_map = {h: k for k, h in spec["cols"]}
    headers: list = []
    seen_keys: set = set()
    for h in rows[0]:
        key = header_map.get(str(h).strip() if h is not None else "", None)
        # Kolom dengan nama berduplikat (mis. kolom "Stok" kedua/G) diabaikan — kemunculan pertama yang dipakai
        if key and key in seen_keys:
            key = None
        if key:
            seen_keys.add(key)
        headers.append(key)
    inserted = updated = skipped = 0
    parts_rows: list[dict] = []
    for r in rows[1:]:
        rec = {headers[i]: (v if v is not None else "") for i, v in enumerate(r) if i < len(headers) and headers[i]}
        if not rec:
            continue
        for k in ("price", "cost", "stock", "min_stock", "km_last"):
            if k in rec:
                try:
                    rec[k] = int(float(rec[k] or 0))
                except (TypeError, ValueError):
                    rec[k] = 0
        if "active" in rec:
            rec["active"] = str(rec["active"]).strip().lower() not in ("false", "0", "tidak", "")
        for k in ("code", "name", "phone", "address", "notes", "barcode", "unit", "rack", "plate", "brand", "model", "year", "color", "chassis_no", "engine_no", "customer_phone", "customer_name"):
            if k in rec:
                rec[k] = str(rec[k]).strip()
        if entity == "customers":
            if not rec.get("name"):
                skipped += 1
                continue
            key = {"phone": rec["phone"]} if rec.get("phone") else {"name": rec["name"]}
            existing = await db.customers.find_one(key)
            if existing:
                await db.customers.update_one({"id": existing["id"]}, {"$set": rec})
                updated += 1
            else:
                seq = await next_counter("customer")
                await db.customers.insert_one({"id": new_id(), "code": rec.get("code") or f"C{seq:04d}", "name": rec["name"], "phone": rec.get("phone", ""),
                                               "address": rec.get("address", ""), "notes": rec.get("notes", ""), "created_at": iso(now()), "deleted_at": None})
                inserted += 1
        elif entity == "parts":
            if not rec.get("code") or not rec.get("name"):
                skipped += 1
                continue
            # Kumpulkan dulu, diproses bulk setelah loop agar file besar tetap cepat
            parts_rows.append(rec)
        elif entity == "services":
            if not rec.get("code") or not rec.get("name"):
                skipped += 1
                continue
            existing = await db.services.find_one({"code": rec["code"], "deleted_at": None})
            if existing:
                await db.services.update_one({"id": existing["id"]}, {"$set": rec})
                updated += 1
            else:
                await db.services.insert_one({"id": new_id(), "active": True, "created_at": iso(now()), "deleted_at": None, **rec})
                inserted += 1
        elif entity == "vehicles":
            plate = (rec.get("plate") or "").upper()
            if not plate:
                skipped += 1
                continue
            cust = None
            if rec.get("customer_phone"):
                cust = await db.customers.find_one({"phone": rec["customer_phone"]})
            if not cust and rec.get("customer_name"):
                cust = await db.customers.find_one({"name": rec["customer_name"]})
            if not cust:
                if not rec.get("customer_name"):
                    skipped += 1
                    continue
                seq = await next_counter("customer")
                cust = {"id": new_id(), "code": f"C{seq:04d}", "name": rec["customer_name"], "phone": rec.get("customer_phone", ""), "address": "", "notes": "",
                        "created_at": iso(now()), "deleted_at": None}
                await db.customers.insert_one(cust)
            vdata = {k: rec.get(k, "") for k in ("brand", "model", "year", "color", "chassis_no", "engine_no")}
            vdata["km_last"] = rec.get("km_last", 0)
            existing = await db.vehicles.find_one({"plate": plate})
            if existing:
                await db.vehicles.update_one({"id": existing["id"]}, {"$set": {**vdata, "customer_id": cust["id"]}})
                updated += 1
            else:
                await db.vehicles.insert_one({"id": new_id(), "plate": plate, "customer_id": cust["id"], **vdata, "last_service_at": None, "created_at": iso(now()), "deleted_at": None})
                inserted += 1
    if entity == "parts" and parts_rows:
        # Satu query untuk semua kode existing, lalu bulk write — tanpa batas jumlah baris
        existing_by_code = {p["code"]: p for p in await db.parts.find({"deleted_at": None}, {"_id": 0}).to_list(None)}
        ops: list = []
        movements: list[dict] = []
        now_iso = iso(now())
        for rec in parts_rows:
            existing = existing_by_code.get(rec["code"])
            if existing:
                if "stock" in rec and rec["stock"] != existing.get("stock"):
                    movements.append({"id": new_id(), "part_id": existing["id"], "part_code": existing["code"], "part_name": rec["name"],
                                      "qty": rec["stock"] - existing.get("stock", 0), "stock_before": existing.get("stock", 0), "stock_after": rec["stock"],
                                      "reason": "IMPORT EXCEL", "transaction_id": None, "transaction_no": None, "username": user["username"], "created_at": now_iso})
                ops.append(UpdateOne({"id": existing["id"]}, {"$set": rec}))
                existing.update(rec)
                updated += 1
            else:
                doc = {"id": new_id(), "active": True, "created_at": now_iso, "deleted_at": None,
                       "barcode": "", "cost": 0, "stock": 0, "min_stock": 5, "unit": "pcs", "price": 0, **rec}
                ops.append(InsertOne(doc))
                existing_by_code[rec["code"]] = doc
                inserted += 1
        for i in range(0, len(ops), 1000):
            await db.parts.bulk_write(ops[i:i + 1000], ordered=True)
        for i in range(0, len(movements), 1000):
            await db.stock_movements.insert_many(movements[i:i + 1000])
    await audit(None, "IMPORT", user, f"Import {entity}: {inserted} baru, {updated} diperbarui, {skipped} dilewati")
    return {"inserted": inserted, "updated": updated, "skipped": skipped}


@api.get("/audit-logs")
async def list_audit(_: OwnerUser, limit: int = 200):
    return await db.audit_logs.find({}, {"_id": 0}).sort("created_at", -1).to_list(limit)


# --------------------------------------------------------------------------- shop profile & logo (Emergent Object Storage)

STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "klinik-suel-motor"
storage_key: Optional[str] = None


def init_storage() -> str:
    global storage_key
    if storage_key:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    resp.raise_for_status()
    return resp.json()


def get_object(path: str) -> tuple[bytes, str]:
    global storage_key
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 503:
        storage_key = None
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": init_storage()}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


class ShopIn(BaseModel):
    name: str
    address: str = ""
    phone: str = ""
    whatsapp: str = ""
    postal_code: str = ""


@api.get("/shop")
async def shop_get():
    return await get_shop()


@api.put("/shop")
async def shop_update(body: ShopIn, user: OwnerUser):
    await db.settings.update_one({"id": "shop"}, {"$set": {**body.model_dump(), "updated_at": iso(now())}}, upsert=True)
    await audit(None, "SHOP_UPDATE", user, "Profil bengkel diperbarui")
    return await get_shop()


@api.post("/shop/logo")
async def shop_logo_upload(user: OwnerUser, file: UploadFile = File(...)):
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran logo maksimal 5 MB")
    ctype = file.content_type or "image/jpeg"
    ext = {"image/png": "png", "image/webp": "webp"}.get(ctype, "jpg")
    path = f"{APP_NAME}/uploads/shop/{uuid.uuid4()}.{ext}"
    try:
        await run_in_threadpool(put_object, path, data, ctype)
    except requests.HTTPError as e:
        code = e.response.status_code if e.response is not None else 500
        if code == 402:
            raise HTTPException(status_code=402, detail="Kuota penyimpanan habis. Tambah saldo untuk mengunggah logo.")
        logger.exception("logo upload failed")
        raise HTTPException(status_code=502, detail="Gagal mengunggah logo ke penyimpanan")
    shop = await get_shop()
    await db.settings.update_one({"id": "shop"}, {"$set": {"logo_path": path, "logo_ctype": ctype, "logo_version": int(shop.get("logo_version", 1)) + 1}}, upsert=True)
    await audit(None, "SHOP_LOGO", user, "Logo bengkel diganti")
    return await get_shop()


@api.delete("/shop/logo")
async def shop_logo_reset(user: OwnerUser):
    shop = await get_shop()
    await db.settings.update_one({"id": "shop"}, {"$set": {"logo_path": None, "logo_ctype": None, "logo_version": int(shop.get("logo_version", 1)) + 1}}, upsert=True)
    await audit(None, "SHOP_LOGO", user, "Logo bengkel dikembalikan ke default")
    return await get_shop()


@api.get("/shop/logo")
async def shop_logo(v: str = ""):
    shop = await get_shop()
    if shop.get("logo_path"):
        try:
            content, ctype = await run_in_threadpool(get_object, shop["logo_path"])
            return Response(content=content, media_type=shop.get("logo_ctype") or ctype, headers={"Cache-Control": "public, max-age=3600"})
        except Exception:
            logger.exception("logo fetch failed, falling back to default")
    return RedirectResponse(shop["logo_url"])


# --------------------------------------------------------------------------- reminders

async def compute_reminders() -> list[dict]:
    today = now().astimezone(timezone(timedelta(hours=7))).date()
    vehicles = await db.vehicles.find({"deleted_at": None, "last_service_at": {"$ne": None}}, {"_id": 0}).to_list(5000)
    cust_ids = list({v["customer_id"] for v in vehicles})
    custs = {c["id"]: c for c in await db.customers.find({"id": {"$in": cust_ids}}, {"_id": 0}).to_list(5000)}
    shop = await get_shop()
    out = []
    for v in vehicles:
        # any active (unpaid) transaction for this vehicle means it's already in the shop
        if await db.service_transactions.count_documents({"vehicle_id": v["id"], "status": {"$in": STATUS_ACTIVE}}):
            continue
        last = await db.service_transactions.find_one({"vehicle_id": v["id"], "status": {"$in": STATUS_PAID}}, {"_id": 0}, sort=[("paid_at", -1)])
        if not last:
            continue
        last_date = parse_date(last.get("paid_at") or last["created_at"])
        due = parse_date(last.get("next_date") or "") or (last_date + timedelta(days=60) if last_date else None)
        if not due:
            continue
        days = (due - today).days
        status = "TERLAMBAT" if days < 0 else "SEGERA" if days <= 7 else "MENDATANG"
        c = custs.get(v["customer_id"], {})
        phone = wa_phone(c.get("phone", ""))
        msg = (f"Halo {c.get('name', 'Pelanggan')}, ini pengingat dari {shop['name']}.\n\n"
               f"Motor {v.get('brand', '')} {v.get('model', '')} ({v['plate']}) dijadwalkan servis berikutnya pada {due.strftime('%d/%m/%Y')}"
               + (f" atau saat KM mencapai {last['next_km']}" if last.get("next_km") else "") + "."
               + (f"\nRekomendasi: {last['next_recommendation']}" if last.get("next_recommendation") else "")
               + f"\n\nSilakan datang ke {shop['name']}, {shop['address']}. Hubungi kami: {shop['phone']}. Terima kasih!")
        from urllib.parse import quote
        sent = await db.whatsapp_logs.find_one({"kind": "REMINDER", "vehicle_id": v["id"], "due_date": due.isoformat()}, {"_id": 0}, sort=[("created_at", -1)])
        out.append({"vehicle_id": v["id"], "plate": v["plate"], "vehicle_name": f"{v.get('brand', '')} {v.get('model', '')}".strip(),
                    "customer_id": v["customer_id"], "customer_name": c.get("name", "-"), "customer_phone": c.get("phone", ""),
                    "last_service_at": last.get("paid_at"), "last_trx_id": last["id"], "next_km": last.get("next_km"), "recommendation": last.get("next_recommendation", ""),
                    "due_date": due.isoformat(), "days_left": days, "status": status, "message": msg,
                    "wa_url": f"https://wa.me/{phone}?text={quote(msg)}" if phone else None, "reminder_sent_at": sent["created_at"] if sent else None})
    out.sort(key=lambda r: r["due_date"])
    return out


@api.get("/reminders")
async def reminders(_: CurrentUser):
    return await compute_reminders()


@api.post("/reminders/{vehicle_id}/sent")
async def reminder_sent(vehicle_id: str, user: CurrentUser, due_date: str = ""):
    v = await db.vehicles.find_one({"id": vehicle_id})
    if not v:
        raise HTTPException(status_code=404, detail="Motor tidak ditemukan")
    await db.whatsapp_logs.insert_one({"id": new_id(), "kind": "REMINDER", "vehicle_id": vehicle_id, "plate": v["plate"], "due_date": due_date,
                                       "status": "TERKIRIM", "sent_by": user["username"], "created_at": iso(now())})
    await audit(None, "REMINDER_SENT", user, f"Pengingat servis dikirim ke pemilik {v['plate']}")
    return {"ok": True}


# --------------------------------------------------------------------------- notifications (owner)

def notif_filter(user: dict) -> dict:
    if user["role"] == "owner":
        return {"$or": [{"roles": "owner"}, {"roles": {"$exists": False}}]}
    return {"roles": user["role"]}


@api.get("/notifications")
async def list_notifications(user: CurrentUser, limit: int = 100):
    docs = await db.notifications.find(notif_filter(user), {"_id": 0}).sort("created_at", -1).to_list(limit)
    for d in docs:
        d["read"] = user["username"] in (d.get("read_by") or []) or d.get("read") is True
    return {"unread": sum(1 for d in docs if not d["read"]), "items": docs}


@api.post("/notifications/{nid}/read")
async def read_notification(nid: str, user: CurrentUser):
    await db.notifications.update_one({"id": nid}, {"$addToSet": {"read_by": user["username"]}, "$set": {"read_at": iso(now())}})
    return {"ok": True}


@api.post("/notifications/read-all")
async def read_all_notifications(user: CurrentUser):
    await db.notifications.update_many(notif_filter(user), {"$addToSet": {"read_by": user["username"]}})
    return {"ok": True}


async def unread_count(user: dict) -> int:
    return await db.notifications.count_documents({**notif_filter(user), "read_by": {"$ne": user["username"]}, "read": {"$ne": True}})


# --------------------------------------------------------------------------- expenses (belanja)

EXPENSE_GROUPS = {
    "BENGKEL": ["Beli Part", "Kasbon Mekanik", "Gaji Mekanik", "Listrik", "Perawatan Gedung", "Tools", "Konsumsi", "Transportasi", "Kerugian", "Kesehatan", "Liburan", "Part Promosi"],
    "KELUARGA": ["Sekolah", "Konsumsi", "Pakaian", "Peralatan Rumah", "Liburan", "Kesehatan", "Saving", "Transportasi", "Komunikasi"],
    "LAINNYA": ["Pinjaman Uang: Orang Tua", "Pinjaman Uang: Bank", "Pinjaman Uang: Pinjol", "Pinjaman Uang: Beli Part",
                "Uang Dipinjam: Orang Lain", "Uang Dipinjam: Kekurangan Bayar Servis", "Uang Dipinjam: Kasbon Mekanik"],
}


class ExpenseIn(BaseModel):
    group: Literal["BENGKEL", "KELUARGA", "LAINNYA"]
    category: str
    amount: int
    note: str = ""
    counterparty: str = ""
    part_id: Optional[str] = None
    qty: int = 0
    date: str = ""  # YYYY-MM-DD optional
    supplier: str = ""
    discount: int = 0
    het: int = 0            # Harga Eceran Tertinggi
    ongkir: int = 0         # ongkos kirim
    unit_price: int = 0     # harga satuan
    payment_method: Literal["CASH", "HUTANG", "TRANSFER"] = "CASH"


@api.get("/expenses/categories")
async def expense_categories(_: CurrentUser):
    return EXPENSE_GROUPS


@api.get("/expenses")
async def list_expenses(_: KasirUser, group: str = "", month: str = ""):
    flt: dict[str, Any] = {"deleted_at": None}
    if group:
        flt["group"] = group
    if month:
        flt["date"] = {"$regex": f"^{month}"}
    return await db.expenses.find(flt, {"_id": 0}).sort("created_at", -1).to_list(2000)


@api.post("/expenses")
async def create_expense(body: ExpenseIn, user: KasirUser):
    if body.amount <= 0:
        raise HTTPException(status_code=400, detail="Nominal harus lebih dari 0")
    if body.category not in EXPENSE_GROUPS[body.group]:
        raise HTTPException(status_code=400, detail="Kategori tidak valid")
    date = body.date or now().astimezone(timezone(timedelta(hours=7))).strftime("%Y-%m-%d")
    doc = {"id": new_id(), **body.model_dump(), "date": date, "part_name": None,
           "flow": ("IN" if body.category.startswith("Pinjaman") else "OUT"), "created_by": user["name"], "created_at": iso(now()), "deleted_at": None}
    if body.group == "BENGKEL" and body.category == "Beli Part" and body.part_id:
        p = await db.parts.find_one({"id": body.part_id, "deleted_at": None})
        if not p:
            raise HTTPException(status_code=404, detail="Part tidak ditemukan")
        if body.qty <= 0:
            raise HTTPException(status_code=400, detail="Qty pembelian harus lebih dari 0")
        await apply_stock_delta(p["id"], body.qty, "PEMBELIAN", None, user)
        doc["part_name"] = p["name"]
    await db.expenses.insert_one(doc)
    await audit(None, "EXPENSE", user, f"Belanja {body.group} · {body.category} Rp{body.amount:,}{' · ' + body.note if body.note else ''}")
    if body.group == "KELUARGA":
        await notify(["owner"], "KASBON OWNER", f"KASBON OWNER: {body.category} Rp {body.amount:,}".replace(",", "."), f"{body.note or '-'}\nDicatat oleh {user['name']} · {date}")
    return clean(doc)


@api.delete("/expenses/{eid}")
async def delete_expense(eid: str, user: KasirUser):
    e = await db.expenses.find_one({"id": eid, "deleted_at": None})
    if not e:
        raise HTTPException(status_code=404, detail="Data belanja tidak ditemukan")
    if user["role"] != "owner" and e.get("created_by") != user["name"]:
        raise HTTPException(status_code=403, detail="Hanya Owner yang dapat menghapus")
    await db.expenses.update_one({"id": eid}, {"$set": {"deleted_at": iso(now())}})
    if e.get("part_id") and e.get("qty"):
        await apply_stock_delta(e["part_id"], -int(e["qty"]), "BATAL PEMBELIAN", None, user)
    await audit(None, "EXPENSE_DELETE", user, f"Hapus belanja {e['category']} Rp{e['amount']:,}")
    return {"ok": True}


class ExpenseEditIn(BaseModel):
    date: Optional[str] = None
    amount: Optional[int] = None
    note: Optional[str] = None
    counterparty: Optional[str] = None
    supplier: Optional[str] = None
    discount: Optional[int] = None
    het: Optional[int] = None
    ongkir: Optional[int] = None
    unit_price: Optional[int] = None
    payment_method: Optional[Literal["CASH", "HUTANG", "TRANSFER"]] = None


@api.put("/expenses/{eid}")
async def edit_expense(eid: str, body: ExpenseEditIn, user: KasirUser):
    e = await db.expenses.find_one({"id": eid, "deleted_at": None})
    if not e:
        raise HTTPException(status_code=404, detail="Data belanja tidak ditemukan")
    if user["role"] != "owner" and e.get("created_by") != user["name"]:
        raise HTTPException(status_code=403, detail="Hanya Owner yang dapat mengubah")
    upd = {k: v for k, v in body.model_dump(exclude_none=True).items()}
    if "date" in upd:
        d = parse_date(upd["date"])
        if not d:
            raise HTTPException(status_code=400, detail="Tanggal tidak valid")
        upd["date"] = d.strftime("%Y-%m-%d")
    if upd:
        await db.expenses.update_one({"id": eid}, {"$set": upd})
    await audit(None, "EXPENSE_EDIT", user, f"Ubah belanja {e['category']}")
    return clean(await db.expenses.find_one({"id": eid}))


@api.get("/expenses/summary")
async def expense_summary(_: KasirUser, month: str = ""):
    month = month or now().astimezone(timezone(timedelta(hours=7))).strftime("%Y-%m")
    rows = await db.expenses.find({"deleted_at": None, "date": {"$regex": f"^{month}"}}, {"_id": 0}).to_list(5000)
    by_group: dict[str, int] = {"BENGKEL": 0, "KELUARGA": 0, "LAINNYA_IN": 0, "LAINNYA_OUT": 0}
    by_cat: dict[str, int] = {}
    for r in rows:
        key = r["group"] if r["group"] != "LAINNYA" else ("LAINNYA_IN" if r.get("flow") == "IN" else "LAINNYA_OUT")
        by_group[key] = by_group.get(key, 0) + r["amount"]
        by_cat[f"{r['group']}: {r['category']}"] = by_cat.get(f"{r['group']}: {r['category']}", 0) + r["amount"]
    pays = await db.payments.find({"date": {"$regex": f"^{month.replace('-', '')}"}, "cancelled": {"$ne": True}}, {"_id": 0, "total": 1}).to_list(10000)
    omzet = sum(p["total"] for p in pays)
    return {"month": month, "omzet": omzet, "belanja_bengkel": by_group["BENGKEL"], "laba_bersih": omzet - by_group["BENGKEL"],
            "kasbon_owner": by_group["KELUARGA"], "pinjaman_masuk": by_group["LAINNYA_IN"], "uang_dipinjam_keluar": by_group["LAINNYA_OUT"],
            "by_category": by_cat, "count": len(rows)}


# --------------------------------------------------------------------------- tools checklist (mekanik)

class ToolIn(BaseModel):
    name: str
    note: str = ""


class ChecklistItemIn(BaseModel):
    tool_id: str
    status: Literal["ADA", "RUSAK", "HILANG"]
    note: str = ""


class ChecklistIn(BaseModel):
    items: list[ChecklistItemIn]
    note: str = ""


@api.get("/tools")
async def list_tools(_: MekanikUser):
    return await db.tools.find({"deleted_at": None}, {"_id": 0}).sort("name", 1).to_list(500)


@api.post("/tools")
async def create_tool(body: ToolIn, user: MekanikUser):
    doc = {"id": new_id(), "name": body.name.strip(), "note": body.note, "added_by": user["name"], "created_at": iso(now()), "deleted_at": None}
    await db.tools.insert_one(doc)
    return clean(doc)


@api.delete("/tools/{tid}")
async def delete_tool(tid: str, user: MekanikUser):
    r = await db.tools.update_one({"id": tid, "deleted_at": None}, {"$set": {"deleted_at": iso(now())}})
    if not r.modified_count:
        raise HTTPException(status_code=404, detail="Peralatan tidak ditemukan")
    return {"ok": True}


async def tools_due_status(user: dict) -> dict:
    last = await db.tool_checklists.find_one({"mechanic_id": user["id"]}, {"_id": 0}, sort=[("created_at", -1)])
    days = None
    if last:
        days = (now() - datetime.fromisoformat(last["created_at"])).days
    due = last is None or (days is not None and days >= 7)
    if due and user["role"] == "mekanik":
        week_ago = iso(now() - timedelta(days=7))
        exists = await db.notifications.find_one({"kind": "CHECKLIST TOOLS", "meta.mechanic_id": user["id"], "created_at": {"$gte": week_ago}})
        if not exists:
            await notify(["mekanik"], "CHECKLIST TOOLS", "Waktunya checklist peralatan mingguan",
                         f"Halo {user['name']}, sudah {days if days is not None else 'lebih dari 7'} hari sejak checklist tools terakhir. Silakan periksa kelengkapan peralatan Anda.",
                         None, {"mechanic_id": user["id"]})
    return {"last_checklist_at": last["created_at"] if last else None, "days_since": days, "due": due,
            "missing": sum(1 for i in (last or {}).get("items", []) if i["status"] != "ADA")}


@api.get("/tools/status")
async def tools_status(user: MekanikUser):
    return await tools_due_status(user)


@api.get("/tools/checklists")
async def list_checklists(user: MekanikUser, all: bool = False):
    flt = {} if (all and user["role"] == "owner") else {"mechanic_id": user["id"]}
    return await db.tool_checklists.find(flt, {"_id": 0}).sort("created_at", -1).to_list(200)


@api.post("/tools/checklists")
async def create_checklist(body: ChecklistIn, user: MekanikUser):
    tools = {t["id"]: t for t in await db.tools.find({"deleted_at": None}, {"_id": 0}).to_list(500)}
    items = [{"tool_id": i.tool_id, "name": tools.get(i.tool_id, {}).get("name", "?"), "status": i.status, "note": i.note} for i in body.items if i.tool_id in tools]
    if not items:
        raise HTTPException(status_code=400, detail="Checklist kosong")
    doc = {"id": new_id(), "mechanic_id": user["id"], "mechanic_name": user["name"], "items": items, "note": body.note,
           "total": len(items), "ok": sum(1 for i in items if i["status"] == "ADA"), "created_at": iso(now())}
    await db.tool_checklists.insert_one(doc)
    missing = [i for i in items if i["status"] != "ADA"]
    if missing:
        await notify(["owner"], "TOOLS BERMASALAH", f"Checklist tools {user['name']}: {len(missing)} peralatan bermasalah",
                     "\n".join(f"- {i['name']}: {i['status']}{' (' + i['note'] + ')' if i['note'] else ''}" for i in missing))
    await audit(None, "TOOLS_CHECKLIST", user, f"Checklist tools: {doc['ok']}/{doc['total']} lengkap")
    return clean(doc)


# --------------------------------------------------------------------------- mechanic report & backup

@api.get("/reports/mechanics")
async def report_mechanics(_: OwnerUser, month: str = ""):
    month = month or now().astimezone(timezone(timedelta(hours=7))).strftime("%Y-%m")
    pays = await db.payments.find({"date": {"$regex": f"^{month.replace('-', '')}"}}, {"_id": 0}).to_list(20000)
    ids = [p["transaction_id"] for p in pays]
    trxs = {t["id"]: t for t in await db.service_transactions.find({"id": {"$in": ids}}, {"_id": 0}).to_list(20000)}
    items = await db.service_items.find({"transaction_id": {"$in": ids}, "deleted_at": None, "approval": "DISETUJUI"}, {"_id": 0}).to_list(100000)
    by_trx: dict[str, dict] = {}
    for i in items:
        b = by_trx.setdefault(i["transaction_id"], {"jasa": 0, "part": 0})
        b[i["kind"]] += i["subtotal"]
    rows: dict[str, dict] = {}
    for p in pays:
        t = trxs.get(p["transaction_id"], {})
        if t.get("status") == "DIBATALKAN":
            continue
        name = t.get("mechanic_name") or "Tanpa mekanik"
        r = rows.setdefault(name, {"mechanic": name, "count": 0, "revenue": 0, "jasa": 0, "part": 0})
        r["count"] += 1
        r["revenue"] += p["total"]
        r["jasa"] += by_trx.get(p["transaction_id"], {}).get("jasa", 0)
        r["part"] += by_trx.get(p["transaction_id"], {}).get("part", 0)
    out = sorted(rows.values(), key=lambda r: -r["revenue"])
    return {"month": month, "rows": out, "total_count": sum(r["count"] for r in out), "total_revenue": sum(r["revenue"] for r in out)}


@api.get("/backup/export")
async def export_backup(_: OwnerUser):
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaksi"
    tcols = ["id", "trx_no", "invoice_no", "queue_no", "date", "status", "customer_name", "customer_phone", "plate", "vehicle_name", "km_in", "complaint_main",
             "mechanic_name", "next_recommendation", "discount", "debt_amount", "debt_status", "debt_due_date", "paid_at", "created_at", "cancel_reason"]
    ws.append(tcols)
    async for t in db.service_transactions.find({}, {"_id": 0}).sort("created_at", 1):
        ws.append([t.get(c, "") for c in tcols])
    sheets = [
        ("Item Transaksi", db.service_items, ["transaction_id", "kind", "source", "approval", "code", "name", "price", "cost", "qty", "subtotal", "note", "added_by", "created_at", "deleted_at"]),
        ("Pembayaran", db.payments, ["transaction_id", "date", "method", "total", "amount_paid", "change", "discount", "reference", "cashier", "paid_at", "due_date"]),
        ("Cicilan Hutang", db.debt_payments, ["transaction_id", "date", "amount", "method", "cashier", "debt_before", "debt_after", "paid_at"]),
        ("Belanja", db.expenses, ["date", "group", "category", "amount", "flow", "note", "counterparty", "supplier", "discount", "het", "ongkir", "unit_price", "payment_method", "part_name", "qty", "created_by", "created_at", "deleted_at"]),
        ("Penjualan Langsung", db.sales, ["sale_no", "date", "outlet_name", "customer_name", "customer_phone", "customer_address", "subtotal", "discount", "total", "total_cost", "total_profit", "method", "amount_paid", "debt_amount", "debt_status", "cashier", "created_at", "deleted_at"]),
        ("Pelanggan", db.customers, ["code", "name", "phone", "address", "dusun", "desa", "kecamatan", "kabupaten", "notes", "created_at", "deleted_at"]),
        ("Motor", db.vehicles, ["plate", "customer_id", "brand", "model", "year", "color", "chassis_no", "engine_no", "km_last", "last_service_at", "deleted_at"]),
        ("Part", db.parts, ["code", "name", "barcode", "price", "cost", "stock", "min_stock", "unit", "rack", "active", "deleted_at"]),
        ("Jasa", db.services, ["code", "name", "price", "active", "deleted_at"]),
        ("Mutasi Stok", db.stock_movements, ["created_at", "part_code", "part_name", "qty", "stock_before", "stock_after", "reason", "transaction_no", "username"]),
        ("Checklist Tools", db.tool_checklists, ["created_at", "mechanic_name", "total", "ok", "note"]),
        ("Audit Log", db.audit_logs, ["created_at", "username", "action", "detail", "transaction_id"]),
    ]
    for title, coll, cols in sheets:
        s = wb.create_sheet(title)
        s.append(cols)
        async for d in coll.find({}, {"_id": 0}).sort("created_at", 1):
            s.append([str(d.get(c, "")) if isinstance(d.get(c), (dict, list)) else d.get(c, "") for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    await db.meta.update_one({"_id": "last_backup"}, {"$set": {"date": wib_today(), "at": iso(now())}}, upsert=True)
    fname = f"suel_backup_{wib_today()}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# --------------------------------------------------------------------------- owner reports & bulk actions

@api.get("/reports/cancellations")
async def report_cancellations(_: OwnerUser, limit: int = 300):
    docs = await db.service_transactions.find({"status": "DIBATALKAN"}, {"_id": 0}).sort("cancelled_at", -1).to_list(limit)
    ids = [d["id"] for d in docs]
    pipeline = [{"$match": {"transaction_id": {"$in": ids}, "deleted_at": None, "approval": "DISETUJUI"}},
                {"$group": {"_id": "$transaction_id", "subtotal": {"$sum": "$subtotal"}}}]
    sums = {r["_id"]: r["subtotal"] for r in await db.service_items.aggregate(pipeline).to_list(1000)}
    logs = {l["transaction_id"]: l for l in await db.transaction_status_logs.find({"transaction_id": {"$in": ids}, "status": "DIBATALKAN"}, {"_id": 0}).to_list(1000)}
    for d in docs:
        d["total"] = max(sums.get(d["id"], 0) - int(d.get("discount", 0)), 0)
        lg = logs.get(d["id"])
        d.setdefault("cancel_reason", (lg or {}).get("note", "").replace("Dibatalkan: ", ""))
        d.setdefault("cancelled_by", (lg or {}).get("username"))
        d.setdefault("cancelled_at", (lg or {}).get("created_at"))
    return {"count": len(docs), "total_value": sum(d["total"] for d in docs), "paid_cancelled": sum(1 for d in docs if d.get("was_paid")), "rows": docs}


@api.get("/export/report/omzet")
async def export_report(_: OwnerUser, mode: Literal["daily", "monthly"] = "daily"):
    rep = await report_omzet(_, mode)
    wb = Workbook()
    ws = wb.active
    ws.title = "Omzet Harian" if mode == "daily" else "Omzet Bulanan"
    ws.append(["Periode", "Jumlah Transaksi", "Omzet Jasa", "Omzet Part", "Diskon", "Total Omzet"])
    for r in rep["rows"]:
        ws.append([r["period"], r["count"], r["jasa"], r["part"], r["diskon"], r["total"]])
    ws.append([])
    ws.append(["TOTAL", sum(r["count"] for r in rep["rows"]), sum(r["jasa"] for r in rep["rows"]), sum(r["part"] for r in rep["rows"]),
               sum(r["diskon"] for r in rep["rows"]), rep["grand_total"]])
    ws2 = wb.create_sheet("Rincian Transaksi")
    ws2.append(["Tanggal Bayar", "No Nota", "No Transaksi", "Pelanggan", "No Polisi", "Mekanik", "Metode", "Total", "Dibayar", "Sisa Hutang", "Kasir"])
    pays = await db.payments.find({}, {"_id": 0}).sort("paid_at", -1).to_list(50000)
    trxs = {t["id"]: t for t in await db.service_transactions.find({"id": {"$in": [p["transaction_id"] for p in pays]}}, {"_id": 0}).to_list(50000)}
    for p in pays:
        t = trxs.get(p["transaction_id"], {})
        ws2.append([p["paid_at"][:19].replace("T", " "), t.get("invoice_no", ""), t.get("trx_no", ""), t.get("customer_name", ""), t.get("plate", ""),
                    t.get("mechanic_name", ""), p["method"], p["total"], p["amount_paid"], t.get("debt_amount", 0) or 0, p.get("cashier", "")])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"suel_laporan_omzet_{mode}_{wib_today()}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@api.post("/parts/delete-all")
async def delete_all_parts(body: OwnerPasswordIn, user: OwnerUser):
    if not await verify_owner_password(body.owner_password):
        raise HTTPException(status_code=400, detail="Password Owner salah")
    r = await db.parts.update_many({"deleted_at": None}, {"$set": {"deleted_at": iso(now())}})
    await audit(None, "PARTS_DELETE_ALL", user, f"Hapus semua data sparepart ({r.modified_count} part)")
    return {"deleted": r.modified_count}


@api.post("/services/delete-all")
async def delete_all_services(body: OwnerPasswordIn, user: OwnerUser):
    if not await verify_owner_password(body.owner_password):
        raise HTTPException(status_code=400, detail="Password Owner salah")
    r = await db.services.update_many({"deleted_at": None}, {"$set": {"deleted_at": iso(now())}})
    await audit(None, "SERVICES_DELETE_ALL", user, f"Hapus semua data jasa ({r.modified_count} jasa)")
    return {"deleted": r.modified_count}


@api.post("/auth/verify-owner")
async def verify_owner(body: OwnerPasswordIn, _: CurrentUser):
    if not await verify_owner_password(body.owner_password):
        raise HTTPException(status_code=400, detail="Password Owner salah")
    return {"ok": True}


# --------------------------------------------------------------------------- outlets & direct part sales

class OutletIn(BaseModel):
    name: str
    address: str = ""
    phone: str = ""


class TransferIn(BaseModel):
    part_id: str
    qty: int
    note: str = ""


class SaleItemIn(BaseModel):
    part_id: str
    qty: int = 1
    price: Optional[int] = None   # harga jual (override master)
    cost: Optional[int] = None    # harga beli (override master)
    discount: int = 0             # diskon per-part (Rp)


class SaleIn(BaseModel):
    outlet_id: str
    items: list[SaleItemIn]
    customer_name: str = "Umum"
    customer_address: str = ""
    customer_phone: str = ""
    customer_dusun: str = ""
    customer_desa: str = ""
    customer_kecamatan: str = ""
    customer_kabupaten: str = ""
    discount: int = 0
    method: Literal["CASH", "TRANSFER", "QRIS", "DEBIT", "KREDIT"] = "CASH"
    amount_paid: int = 0
    due_date: str = ""
    sale_date: str = ""
    note: str = ""


async def ensure_main_outlet() -> dict:
    main = await db.outlets.find_one({"is_main": True, "deleted_at": None}, {"_id": 0})
    if not main:
        main = {"id": new_id(), "name": "Bengkel Pusat", "address": "", "phone": "", "is_main": True, "created_at": iso(now()), "deleted_at": None}
        await db.outlets.insert_one(dict(main))
    return main


@api.get("/outlets")
async def list_outlets(_: SalesUser):
    await ensure_main_outlet()
    return await db.outlets.find({"deleted_at": None}, {"_id": 0}).sort("created_at", 1).to_list(100)


@api.post("/outlets")
async def create_outlet(body: OutletIn, user: OwnerUser):
    doc = {"id": new_id(), **body.model_dump(), "is_main": False, "created_at": iso(now()), "deleted_at": None}
    await db.outlets.insert_one(doc)
    await audit(None, "OUTLET_CREATE", user, f"Outlet baru {body.name}")
    return clean(doc)


@api.get("/outlets/{oid}/stock")
async def outlet_stock(oid: str, _: SalesUser, q: str = ""):
    outlet = await db.outlets.find_one({"id": oid, "deleted_at": None})
    if not outlet:
        raise HTTPException(status_code=404, detail="Outlet tidak ditemukan")
    flt: dict[str, Any] = {"deleted_at": None}
    if q.strip():
        flt["$or"] = [{"code": regex(q)}, {"name": regex(q)}, {"barcode": regex(q)}, {"rack": regex(q)}]
    parts = await db.parts.find(flt, {"_id": 0}).sort("name", 1).to_list(1000)
    if outlet.get("is_main"):
        for p in parts:
            p["outlet_stock"] = p["stock"]
        return parts
    os_map = {s["part_id"]: s["stock"] for s in await db.outlet_stocks.find({"outlet_id": oid}, {"_id": 0}).to_list(5000)}
    for p in parts:
        p["outlet_stock"] = os_map.get(p["id"], 0)
    return parts


@api.post("/outlets/{oid}/transfer")
async def transfer_stock(oid: str, body: TransferIn, user: SalesUser):
    outlet = await db.outlets.find_one({"id": oid, "deleted_at": None})
    if not outlet or outlet.get("is_main"):
        raise HTTPException(status_code=400, detail="Pilih outlet cabang tujuan")
    part = await db.parts.find_one({"id": body.part_id, "deleted_at": None})
    if not part:
        raise HTTPException(status_code=404, detail="Part tidak ditemukan")
    if body.qty > 0 and part["stock"] < body.qty:
        raise HTTPException(status_code=400, detail=f"Stok pusat tidak cukup (tersedia {part['stock']})")
    cur = await db.outlet_stocks.find_one({"outlet_id": oid, "part_id": part["id"]})
    if body.qty < 0 and (cur or {}).get("stock", 0) < -body.qty:
        raise HTTPException(status_code=400, detail="Stok outlet tidak cukup untuk dikembalikan")
    await apply_stock_delta(part["id"], -body.qty, f"TRANSFER KE {outlet['name'].upper()}" if body.qty > 0 else f"RETUR DARI {outlet['name'].upper()}", None, user)
    await db.outlet_stocks.update_one({"outlet_id": oid, "part_id": part["id"]}, {"$inc": {"stock": body.qty}, "$set": {"part_code": part["code"], "part_name": part["name"]}}, upsert=True)
    await audit(None, "STOCK_TRANSFER", user, f"Transfer {part['name']} {body.qty:+d} → {outlet['name']}. {body.note}")
    return {"ok": True}


@api.post("/sales")
async def create_sale(body: SaleIn, user: SalesUser):
    outlet = await db.outlets.find_one({"id": body.outlet_id, "deleted_at": None})
    if not outlet:
        raise HTTPException(status_code=404, detail="Outlet tidak ditemukan")
    if not body.items:
        raise HTTPException(status_code=400, detail="Belum ada part yang dijual")
    items = []
    for it in body.items:
        p = await db.parts.find_one({"id": it.part_id, "deleted_at": None})
        if not p or it.qty <= 0:
            raise HTTPException(status_code=400, detail="Part tidak valid")
        avail = p["stock"] if outlet.get("is_main") else ((await db.outlet_stocks.find_one({"outlet_id": outlet["id"], "part_id": p["id"]})) or {}).get("stock", 0)
        if avail < it.qty:
            raise HTTPException(status_code=400, detail=f"STOK TIDAK CUKUP: {p['name']} (tersedia {avail} di {outlet['name']})")
        price = int(it.price) if it.price is not None else int(p.get("price", 0))
        cost = int(it.cost) if it.cost is not None else int(p.get("cost", 0))
        item_disc = max(int(it.discount or 0), 0)
        subtotal = max(price * it.qty - item_disc, 0)
        profit = subtotal - cost * it.qty
        items.append({"part_id": p["id"], "code": p["code"], "name": p["name"], "rack": p.get("rack", ""), "price": price, "cost": cost,
                      "qty": it.qty, "discount": item_disc, "subtotal": subtotal, "profit": profit})
    subtotal = sum(i["subtotal"] for i in items)
    total = max(subtotal - body.discount, 0)
    total_cost = sum(i["cost"] * i["qty"] for i in items)
    total_profit = total - total_cost
    is_credit = body.method == "KREDIT"
    if not is_credit and body.method != "CASH":
        amount_paid = total
    else:
        amount_paid = int(body.amount_paid or 0)
    if not is_credit and amount_paid < total:
        raise HTTPException(status_code=400, detail="Jumlah dibayar kurang dari total")
    debt = max(total - amount_paid, 0) if is_credit else 0
    today = wib_today()
    sale_date = (body.sale_date or "").strip() or today
    seq = await next_counter("sale")
    sale = {"id": new_id(), "sale_no": f"PJ-{sale_date}-{seq:04d}", "outlet_id": outlet["id"], "outlet_name": outlet["name"], "customer_name": body.customer_name or "Umum",
            "customer_address": body.customer_address, "customer_phone": body.customer_phone, "items": items, "subtotal": subtotal, "discount": body.discount,
            "customer_dusun": body.customer_dusun, "customer_desa": body.customer_desa, "customer_kecamatan": body.customer_kecamatan, "customer_kabupaten": body.customer_kabupaten,
            "total": total, "total_cost": total_cost, "total_profit": total_profit, "method": body.method, "amount_paid": amount_paid,
            "change": max(amount_paid - total, 0) if body.method == "CASH" else 0, "debt_amount": debt,
            "debt_status": ("BELUM LUNAS" if debt > 0 else ("LUNAS" if is_credit else None)), "debt_due_date": body.due_date if debt > 0 else "", "installments": [],
            "note": body.note, "date": sale_date, "cashier": user["name"], "status": "SELESAI", "created_at": iso(now()), "deleted_at": None}
    for i in items:
        if outlet.get("is_main"):
            await apply_stock_delta(i["part_id"], -i["qty"], "PENJUALAN LANGSUNG", {"id": sale["id"], "trx_no": sale["sale_no"]}, user)
        else:
            r = await db.outlet_stocks.find_one_and_update({"outlet_id": outlet["id"], "part_id": i["part_id"]}, {"$inc": {"stock": -i["qty"]}}, return_document=True)
            await db.stock_movements.insert_one({"id": new_id(), "part_id": i["part_id"], "part_code": i["code"], "part_name": i["name"], "qty": -i["qty"],
                                                 "stock_before": r["stock"] + i["qty"], "stock_after": r["stock"], "reason": f"PENJUALAN {outlet['name'].upper()}",
                                                 "transaction_id": sale["id"], "transaction_no": sale["sale_no"], "username": user["username"], "created_at": iso(now())})
    await db.sales.insert_one(sale)
    await db.payments.insert_one({"id": new_id(), "transaction_id": sale["id"], "sale": True, "outlet_id": outlet["id"], "method": body.method, "total": total,
                                  "amount_paid": amount_paid, "change": sale["change"], "discount": body.discount, "reference": "", "note": body.note,
                                  "cashier": user["name"], "paid_at": iso(now()), "date": sale_date, "due_date": sale["debt_due_date"], "installments": []})
    await audit(None, "SALE", user, f"Penjualan part {sale['sale_no']} di {outlet['name']} Rp{total:,} (laba Rp{total_profit:,})" + (f" [hutang Rp{debt:,}]" if debt else ""))
    if debt > 0:
        kind = "HUTANG PENUH" if amount_paid == 0 else "HUTANG SEBAGIAN"
        await notify_owner("HUTANG", f"{kind} (Penjualan Part): {sale['customer_name']}",
                           f"Faktur {sale['sale_no']} · {outlet['name']}\nNominal hutang: Rp {debt:,}".replace(",", ".") + f"\nJatuh tempo: {body.due_date or '-'}\nKasir: {user['name']}",
                           None, {"customer_name": sale["customer_name"], "amount": debt, "due_date": body.due_date, "invoice_no": sale["sale_no"], "sale_id": sale["id"]})
    return clean(sale)


@api.post("/sales/{sid}/pay-debt")
async def pay_sale_debt(sid: str, body: DebtPayIn, user: KasirUser):
    sale = await db.sales.find_one({"id": sid})
    if not sale:
        raise HTTPException(status_code=404, detail="Faktur tidak ditemukan")
    debt = int(sale.get("debt_amount") or 0)
    if debt <= 0:
        raise HTTPException(status_code=400, detail="Faktur tidak memiliki hutang")
    if body.amount <= 0:
        raise HTTPException(status_code=400, detail="Nominal pembayaran tidak valid")
    amount = min(body.amount, debt)
    new_debt = debt - amount
    inst = {"id": new_id(), "amount": amount, "method": body.method, "reference": body.reference, "note": body.note, "cashier": user["name"], "paid_at": iso(now()), "date": wib_today()}
    await db.sales.update_one({"id": sid}, {"$inc": {"amount_paid": amount}, "$set": {"debt_amount": new_debt, "debt_status": "LUNAS" if new_debt == 0 else "BELUM LUNAS"}, "$push": {"installments": inst}})
    await db.payments.update_one({"transaction_id": sid}, {"$inc": {"amount_paid": amount}, "$push": {"installments": inst}})
    await db.debt_payments.insert_one({**inst, "transaction_id": sid, "sale": True, "debt_before": debt, "debt_after": new_debt})
    await audit(None, "SALE_DEBT_PAYMENT", user, f"Pembayaran hutang penjualan {sale.get('sale_no')} {body.method} Rp{amount:,}. Sisa Rp{new_debt:,}")
    if new_debt == 0:
        await notify_owner("HUTANG LUNAS", f"HUTANG LUNAS (Penjualan): {sale.get('customer_name')}", f"Faktur {sale.get('sale_no')}\nHutang telah dilunasi (Rp {amount:,}).".replace(",", "."), None, {"sale_id": sid})
    return clean(await db.sales.find_one({"id": sid}))


@api.get("/sales")
async def list_sales(_: SalesUser, outlet_id: str = "", q: str = "", limit: int = 200):
    flt: dict[str, Any] = {"deleted_at": None}
    if outlet_id:
        flt["outlet_id"] = outlet_id
    if q.strip():
        flt["$or"] = [{"sale_no": regex(q)}, {"customer_name": regex(q)}, {"customer_phone": regex(q)}, {"items.name": regex(q)}]
    return await db.sales.find(flt, {"_id": 0}).sort("created_at", -1).to_list(limit)


@api.get("/sales/{sid}")
async def sale_detail(sid: str, _: SalesUser):
    s = await db.sales.find_one({"id": sid}, {"_id": 0})
    if not s:
        raise HTTPException(status_code=404, detail="Faktur tidak ditemukan")
    return {"sale": s, "shop": await get_shop()}


@api.post("/sales/{sid}/whatsapp")
async def sale_whatsapp(sid: str, user: SalesUser):
    s = await db.sales.find_one({"id": sid}, {"_id": 0})
    if not s:
        raise HTTPException(status_code=404, detail="Faktur tidak ditemukan")
    shop = await get_shop()
    from urllib.parse import quote
    lines = "\n".join(f"- {i['name']} x{i['qty']} = Rp {i['subtotal']:,}".replace(",", ".") for i in s["items"])
    msg = (f"Terima kasih telah berbelanja di {shop['name']} ({s['outlet_name']}).\n\nNo Faktur: {s['sale_no']}\nTanggal: {ymd_display(s['date'])}\n"
           f"Pelanggan: {s['customer_name']}\n{lines}\nTotal: Rp {s['total']:,}".replace(",", ".") + f"\n\nTerima kasih!")
    phone = wa_phone(s.get("customer_phone", ""))
    await db.whatsapp_logs.insert_one({"id": new_id(), "kind": "SALE", "sale_id": sid, "phone": phone, "message": msg, "status": "TERKIRIM" if phone else "GAGAL", "sent_by": user["username"], "created_at": iso(now())})
    return {"status": "TERKIRIM" if phone else "GAGAL", "message": msg, "url": f"https://wa.me/{phone}?text={quote(msg)}" if phone else None}


def ymd_display(ymd: str) -> str:
    return f"{ymd[6:8]}/{ymd[4:6]}/{ymd[0:4]}" if ymd and len(ymd) == 8 else ymd


# --------------------------------------------------------------------------- edit tanggal transaksi (owner)

class DateEditIn(BaseModel):
    date: str  # DD/MM/YYYY atau YYYY-MM-DD


def shift_iso_to_date(iso_ts: Optional[str], new_date) -> str:
    """Ganti komponen tanggal pada timestamp ISO, pertahankan jam:menit:detik. Fallback jam 12:00 WIB."""
    wib = timezone(timedelta(hours=7))
    try:
        dt = datetime.fromisoformat(iso_ts).astimezone(wib) if iso_ts else None
    except (ValueError, TypeError):
        dt = None
    if dt is None:
        dt = datetime(new_date.year, new_date.month, new_date.day, 12, 0, tzinfo=wib)
    else:
        dt = dt.replace(year=new_date.year, month=new_date.month, day=new_date.day)
    return dt.astimezone(timezone.utc).isoformat()


@api.put("/transactions/{tid}/date")
async def edit_transaction_date(tid: str, body: DateEditIn, user: OwnerUser):
    trx = await get_trx(tid)
    d = parse_date(body.date)
    if not d:
        raise HTTPException(status_code=400, detail="Format tanggal tidak valid")
    ymd = d.strftime("%Y%m%d")
    upd = {"date": ymd, "updated_at": iso(now())}
    if trx.get("paid_at"):
        upd["paid_at"] = shift_iso_to_date(trx.get("paid_at"), d)
    await db.service_transactions.update_one({"id": tid}, {"$set": upd})
    pay = await db.payments.find_one({"transaction_id": tid})
    if pay:
        await db.payments.update_one({"transaction_id": tid}, {"$set": {"date": ymd, "paid_at": shift_iso_to_date(pay.get("paid_at"), d)}})
    await db.service_history.update_one({"transaction_id": tid}, {"$set": {"date": ymd}})
    await audit(tid, "DATE_EDIT", user, f"Ubah tanggal transaksi menjadi {d.strftime('%d/%m/%Y')}")
    return await build_detail(await get_trx(tid))


@api.put("/sales/{sid}/date")
async def edit_sale_date(sid: str, body: DateEditIn, user: SalesUser):
    if user["role"] not in ("owner",):
        raise HTTPException(status_code=403, detail="Hanya Owner yang dapat mengubah tanggal")
    sale = await db.sales.find_one({"id": sid})
    if not sale:
        raise HTTPException(status_code=404, detail="Faktur tidak ditemukan")
    d = parse_date(body.date)
    if not d:
        raise HTTPException(status_code=400, detail="Format tanggal tidak valid")
    ymd = d.strftime("%Y%m%d")
    await db.sales.update_one({"id": sid}, {"$set": {"date": ymd}})
    pay = await db.payments.find_one({"transaction_id": sid})
    if pay:
        await db.payments.update_one({"transaction_id": sid}, {"$set": {"date": ymd, "paid_at": shift_iso_to_date(pay.get("paid_at"), d)}})
    await audit(None, "SALE_DATE_EDIT", user, f"Ubah tanggal penjualan {sale.get('sale_no')} menjadi {d.strftime('%d/%m/%Y')}")
    return clean(await db.sales.find_one({"id": sid}))


# --------------------------------------------------------------------------- laporan penjualan servis / jualan / belanja

def month_or_now(month: str) -> str:
    return month or now().astimezone(timezone(timedelta(hours=7))).strftime("%Y-%m")


def ymd_range_filter(month: str, dfrom: str, dto: str) -> dict:
    """Filter untuk field tanggal berformat YYYYMMDD."""
    if dfrom or dto:
        r: dict = {}
        if dfrom:
            r["$gte"] = dfrom.replace("-", "")
        if dto:
            r["$lte"] = dto.replace("-", "")
        return r
    return {"$regex": f"^{month.replace('-', '')}"}


def dash_range_filter(month: str, dfrom: str, dto: str) -> dict:
    """Filter untuk field tanggal berformat YYYY-MM-DD."""
    if dfrom or dto:
        r: dict = {}
        if dfrom:
            r["$gte"] = dfrom
        if dto:
            r["$lte"] = dto
        return r
    return {"$regex": f"^{month}"}


async def build_service_sales(month: str, dfrom: str = "", dto: str = "") -> dict:
    trxs = await db.service_transactions.find({"date": ymd_range_filter(month, dfrom, dto), "status": {"$in": STATUS_PAID}, "deleted_at": None}, {"_id": 0}).sort("date", 1).to_list(20000)
    ids = [t["id"] for t in trxs]
    all_items = await db.service_items.find({"transaction_id": {"$in": ids}, "deleted_at": None, "approval": "DISETUJUI"}, {"_id": 0}).to_list(100000)
    part_ids = [i.get("ref_id") for i in all_items if i["kind"] == "part" and i.get("ref_id")]
    part_codes = [i.get("code") for i in all_items if i["kind"] == "part" and i.get("code")]
    part_cost = {p["id"]: int(p.get("cost", 0) or 0) for p in await db.parts.find({"id": {"$in": part_ids}}, {"_id": 0, "id": 1, "cost": 1}).to_list(100000)}
    cost_by_code = {p["code"]: int(p.get("cost", 0) or 0) for p in await db.parts.find({"code": {"$in": part_codes}}, {"_id": 0, "code": 1, "cost": 1}).to_list(100000)}
    by_trx: dict[str, list] = {}
    for i in all_items:
        by_trx.setdefault(i["transaction_id"], []).append(i)
    rows = []
    for t in trxs:
        its = by_trx.get(t["id"], [])
        jasa = [{"name": i["name"], "price": i["price"], "qty": i["qty"], "subtotal": i["subtotal"]} for i in its if i["kind"] == "jasa"]
        parts = []
        for i in its:
            if i["kind"] != "part":
                continue
            cost = int(i.get("cost") or 0) or part_cost.get(i.get("ref_id"), 0) or cost_by_code.get(i.get("code"), 0)
            profit = i["subtotal"] - cost * i["qty"]
            parts.append({"name": i["name"], "code": i.get("code", ""), "price": i["price"], "cost": cost, "qty": i["qty"], "subtotal": i["subtotal"], "profit": profit})
        total_jasa = sum(j["subtotal"] for j in jasa)
        total_part = sum(p["subtotal"] for p in parts)
        total_part_cost = sum(p["cost"] * p["qty"] for p in parts)
        rows.append({
            "id": t["id"], "date": t.get("date"), "trx_no": t.get("trx_no"), "invoice_no": t.get("invoice_no"),
            "customer_name": t.get("customer_name"), "customer_phone": t.get("customer_phone"),
            "plate": t.get("plate"), "vehicle_name": t.get("vehicle_name"), "mechanic_name": t.get("mechanic_name"),
            "complaint_main": t.get("complaint_main"), "next_recommendation": t.get("next_recommendation", ""),
            "jasa_items": jasa, "part_items": parts, "total_jasa": total_jasa, "total_part": total_part,
            "total_part_cost": total_part_cost, "total_part_profit": total_part - total_part_cost,
            "discount": int(t.get("discount", 0) or 0), "total": max(total_jasa + total_part - int(t.get("discount", 0) or 0), 0),
        })
    return {
        "month": month, "rows": rows, "count": len(rows),
        "total_jasa": sum(r["total_jasa"] for r in rows), "total_part": sum(r["total_part"] for r in rows),
        "total_part_profit": sum(r["total_part_profit"] for r in rows), "grand_total": sum(r["total"] for r in rows),
    }


@api.get("/reports/service-sales")
async def report_service_sales(_: OwnerUser, month: str = "", date_from: str = "", date_to: str = ""):
    return await build_service_sales(month_or_now(month), date_from, date_to)


async def build_direct_sales(month: str, dfrom: str = "", dto: str = "") -> dict:
    sales = await db.sales.find({"date": ymd_range_filter(month, dfrom, dto), "deleted_at": None}, {"_id": 0}).sort("date", 1).to_list(20000)
    rows = []
    for s in sales:
        rows.append({
            "id": s["id"], "date": s.get("date"), "sale_no": s.get("sale_no"), "outlet_name": s.get("outlet_name"),
            "customer_name": s.get("customer_name"), "customer_phone": s.get("customer_phone"), "customer_address": s.get("customer_address"),
            "items": s.get("items", []), "subtotal": s.get("subtotal", 0), "discount": s.get("discount", 0),
            "total": s.get("total", 0), "total_cost": s.get("total_cost", 0), "total_profit": s.get("total_profit", 0),
            "method": s.get("method"), "debt_amount": s.get("debt_amount", 0), "debt_status": s.get("debt_status"),
        })
    return {
        "month": month, "rows": rows, "count": len(rows),
        "grand_total": sum(r["total"] for r in rows), "total_profit": sum(r["total_profit"] for r in rows),
    }


@api.get("/reports/direct-sales")
async def report_direct_sales(_: OwnerUser, month: str = "", date_from: str = "", date_to: str = ""):
    return await build_direct_sales(month_or_now(month), date_from, date_to)


async def build_purchases(month: str, dfrom: str = "", dto: str = "") -> dict:
    rows = await db.expenses.find({"deleted_at": None, "date": dash_range_filter(month, dfrom, dto)}, {"_id": 0}).sort("date", 1).to_list(20000)
    parts, kasbon_mekanik, kasbon_owner, operasional, lainnya = [], [], [], [], []
    by_method = {"CASH": 0, "HUTANG": 0, "TRANSFER": 0}
    for r in rows:
        pm = r.get("payment_method") or "CASH"
        if pm in by_method:
            by_method[pm] += int(r.get("amount", 0) or 0)
        rec = {
            "id": r["id"], "date": r.get("date"), "category": r.get("category"), "amount": int(r.get("amount", 0) or 0),
            "note": r.get("note", ""), "supplier": r.get("supplier", "") or r.get("counterparty", ""), "part_name": r.get("part_name"),
            "qty": r.get("qty", 0), "discount": int(r.get("discount", 0) or 0), "het": int(r.get("het", 0) or 0),
            "ongkir": int(r.get("ongkir", 0) or 0), "unit_price": int(r.get("unit_price", 0) or 0), "payment_method": pm, "flow": r.get("flow", "OUT"),
        }
        if r["group"] == "KELUARGA":
            kasbon_owner.append(rec)
        elif r["group"] == "LAINNYA":
            lainnya.append(rec)
        elif r.get("category") == "Beli Part":
            parts.append(rec)
        elif r.get("category") in ("Kasbon Mekanik", "Gaji Mekanik"):
            kasbon_mekanik.append(rec)
        else:
            operasional.append(rec)
    tot = lambda xs: sum(x["amount"] for x in xs)  # noqa: E731
    return {
        "month": month, "parts": parts, "kasbon_mekanik": kasbon_mekanik, "kasbon_owner": kasbon_owner, "operasional": operasional, "lainnya": lainnya,
        "total_parts": tot(parts), "total_kasbon_mekanik": tot(kasbon_mekanik), "total_kasbon_owner": tot(kasbon_owner),
        "total_operasional": tot(operasional), "total_lainnya": tot(lainnya), "by_method": by_method,
        "grand_total": tot(parts) + tot(kasbon_mekanik) + tot(kasbon_owner) + tot(operasional),
    }


@api.get("/reports/purchases")
async def report_purchases(_: KasirUser, month: str = "", date_from: str = "", date_to: str = ""):
    return await build_purchases(month_or_now(month), date_from, date_to)


def xlsx_response(wb, fname: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@api.get("/export/report/service-sales")
async def export_service_sales(_: OwnerUser, month: str = "", date_from: str = "", date_to: str = ""):
    rep = await build_service_sales(month_or_now(month), date_from, date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "Penjualan Servis"
    ws.append(["Tanggal", "No Nota", "Pelanggan", "No HP", "No Polisi", "Motor", "Mekanik", "Saran Servis",
               "Jenis", "Item", "Qty", "Harga Jual", "Harga Beli", "Subtotal", "Profit"])
    for r in rep["rows"]:
        tgl = ymd_display(r["date"] or "")
        for j in r["jasa_items"]:
            ws.append([tgl, r["invoice_no"] or r["trx_no"], r["customer_name"], r["customer_phone"], r["plate"], r["vehicle_name"], r["mechanic_name"], r["next_recommendation"],
                       "JASA", j["name"], j["qty"], j["price"], "", j["subtotal"], ""])
        for p in r["part_items"]:
            ws.append([tgl, r["invoice_no"] or r["trx_no"], r["customer_name"], r["customer_phone"], r["plate"], r["vehicle_name"], r["mechanic_name"], r["next_recommendation"],
                       "PART", p["name"], p["qty"], p["price"], p["cost"], p["subtotal"], p["profit"]])
        ws.append(["", "", "", "", "", "", "", f"TOTAL NOTA {r['invoice_no'] or r['trx_no']}", "", "", "", "", "", r["total"], r["total_part_profit"]])
        ws.append([])
    ws.append(["GRAND TOTAL", "", "", "", "", "", "", "", "", "", "", "", "", rep["grand_total"], rep["total_part_profit"]])
    return xlsx_response(wb, f"suel_penjualan_servis_{month_or_now(month)}.xlsx")


@api.get("/export/report/direct-sales")
async def export_direct_sales(_: OwnerUser, month: str = "", date_from: str = "", date_to: str = ""):
    rep = await build_direct_sales(month_or_now(month), date_from, date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "Jualan Langsung"
    ws.append(["Tanggal", "No Faktur", "Outlet", "Pelanggan", "No HP", "Alamat", "Item", "Qty", "Harga Jual", "Harga Beli", "Diskon", "Subtotal", "Profit", "Metode"])
    for r in rep["rows"]:
        tgl = ymd_display(r["date"] or "")
        for it in r["items"]:
            ws.append([tgl, r["sale_no"], r["outlet_name"], r["customer_name"], r["customer_phone"], r["customer_address"],
                       it.get("name"), it.get("qty"), it.get("price"), it.get("cost"), it.get("discount", 0), it.get("subtotal"), it.get("profit"), r["method"]])
        ws.append(["", "", "", "", "", "", f"TOTAL {r['sale_no']}", "", "", "", "", r["total"], r["total_profit"], ""])
        ws.append([])
    ws.append(["GRAND TOTAL", "", "", "", "", "", "", "", "", "", "", rep["grand_total"], rep["total_profit"], ""])
    return xlsx_response(wb, f"suel_jualan_langsung_{month_or_now(month)}.xlsx")


@api.get("/export/report/purchases")
async def export_purchases(_: KasirUser, month: str = "", date_from: str = "", date_to: str = ""):
    rep = await build_purchases(month_or_now(month), date_from, date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pembelian Part"
    ws.append(["Tanggal", "Nama Part", "Supplier", "Qty", "Harga Satuan", "HET", "Diskon", "Ongkir", "Total", "Metode", "Catatan"])
    for r in rep["parts"]:
        ws.append([ymd_display(r["date"] or ""), r["part_name"] or r["category"], r["supplier"], r["qty"], r["unit_price"], r["het"], r["discount"], r["ongkir"], r["amount"], r["payment_method"], r["note"]])
    ws.append(["TOTAL PEMBELIAN PART", "", "", "", "", "", "", "", rep["total_parts"], "", ""])
    ws2 = wb.create_sheet("Kasbon & Operasional")
    ws2.append(["Tanggal", "Jenis", "Kategori", "Penerima/Supplier", "Total", "Metode", "Catatan"])
    for label, items in [("KASBON MEKANIK", rep["kasbon_mekanik"]), ("KASBON OWNER", rep["kasbon_owner"]), ("BIAYA OPERASIONAL", rep["operasional"]), ("LAIN-LAIN", rep["lainnya"])]:
        for r in items:
            ws2.append([ymd_display(r["date"] or ""), label, r["category"], r["supplier"], r["amount"], r["payment_method"], r["note"]])
    ws2.append([])
    ws2.append(["Total Kasbon Mekanik", "", "", "", rep["total_kasbon_mekanik"], "", ""])
    ws2.append(["Total Kasbon Owner", "", "", "", rep["total_kasbon_owner"], "", ""])
    ws2.append(["Total Biaya Operasional", "", "", "", rep["total_operasional"], "", ""])
    ws2.append(["CASH / HUTANG / TRANSFER", "", "", "", f"{rep['by_method']['CASH']} / {rep['by_method']['HUTANG']} / {rep['by_method']['TRANSFER']}", "", ""])
    return xlsx_response(wb, f"suel_belanja_{month_or_now(month)}.xlsx")


# --------------------------------------------------------------------------- restore / import backup

def _rows(wb, name: str) -> list[dict]:
    if name not in wb.sheetnames:
        return []
    r = list(wb[name].iter_rows(values_only=True))
    if not r:
        return []
    cols = [str(c).strip() if c is not None else "" for c in r[0]]
    return [{cols[i]: v for i, v in enumerate(row) if i < len(cols) and cols[i]} for row in r[1:]]


def _i(v) -> int:
    try:
        return int(float(v or 0))
    except (TypeError, ValueError):
        return 0


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _dt_of(s):
    try:
        return datetime.fromisoformat(_s(s))
    except (ValueError, TypeError):
        return None


@api.post("/backup/import")
async def import_backup(user: OwnerUser, file: UploadFile = File(...)):
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="File bukan Excel (.xlsx) yang valid")
    # Buang data hasil impor sebelumnya agar idempotent (tidak menyentuh data yang diinput manual)
    for coll in ("service_transactions", "service_items", "payments", "debt_payments", "service_history", "expenses", "sales", "stock_movements"):
        await db[coll].delete_many({"_import": True})
    now_iso = iso(now())
    stats = {"customers": 0, "vehicles": 0, "parts": 0, "services": 0, "transactions": 0, "payments": 0, "expenses": 0, "sales": 0}

    # ---- Pelanggan
    cust_by_phone: dict[str, dict] = {}
    cust_by_name: dict[str, dict] = {}
    for r in _rows(wb, "Pelanggan"):
        name = _s(r.get("name"))
        if not name:
            continue
        phone = _s(r.get("phone"))
        existing = None
        if phone:
            existing = await db.customers.find_one({"phone": phone, "deleted_at": None})
        if not existing:
            existing = await db.customers.find_one({"name": name, "deleted_at": None})
        data = {"name": name, "phone": phone, "address": _s(r.get("address")), "notes": _s(r.get("notes")),
                "dusun": _s(r.get("dusun")), "desa": _s(r.get("desa")), "kecamatan": _s(r.get("kecamatan")), "kabupaten": _s(r.get("kabupaten"))}
        if existing:
            await db.customers.update_one({"id": existing["id"]}, {"$set": data})
            cid = existing["id"]
        else:
            seq = await next_counter("customer")
            doc = {"id": new_id(), "code": _s(r.get("code")) or f"C{seq:04d}", **data, "created_at": _s(r.get("created_at")) or now_iso, "deleted_at": r.get("deleted_at") or None}
            await db.customers.insert_one(doc)
            cid = doc["id"]
            stats["customers"] += 1
        cust = await db.customers.find_one({"id": cid}, {"_id": 0})
        if phone:
            cust_by_phone[phone] = cust
        cust_by_name[name.lower()] = cust

    # plate -> pemilik, dari sheet Transaksi (Motor tidak punya id pelanggan yang cocok)
    plate_owner: dict[str, tuple] = {}
    for t in _rows(wb, "Transaksi"):
        pl = _s(t.get("plate")).upper()
        if pl:
            plate_owner[pl] = (_s(t.get("customer_name")), _s(t.get("customer_phone")))

    async def fallback_customer() -> dict:
        c = await db.customers.find_one({"name": "Pelanggan Umum"}, {"_id": 0})
        if not c:
            seq = await next_counter("customer")
            c = {"id": new_id(), "code": f"C{seq:04d}", "name": "Pelanggan Umum", "phone": "", "address": "", "notes": "Dibuat saat restore", "created_at": now_iso, "deleted_at": None}
            await db.customers.insert_one(dict(c))
        return c

    # ---- Motor
    for r in _rows(wb, "Motor"):
        plate = _s(r.get("plate")).upper()
        if not plate:
            continue
        cust = None
        owner = plate_owner.get(plate)
        if owner:
            oname, ophone = owner
            cust = (cust_by_phone.get(ophone) if ophone else None) or cust_by_name.get(oname.lower())
        if not cust:
            cust = await fallback_customer()
        vdata = {k: _s(r.get(k)) for k in ("brand", "model", "year", "color", "chassis_no", "engine_no")}
        vdata["km_last"] = _i(r.get("km_last"))
        existing = await db.vehicles.find_one({"plate": plate})
        if existing:
            await db.vehicles.update_one({"id": existing["id"]}, {"$set": {**vdata, "customer_id": cust["id"]}})
        else:
            await db.vehicles.insert_one({"id": new_id(), "plate": plate, "customer_id": cust["id"], **vdata,
                                          "last_service_at": r.get("last_service_at") or None, "created_at": now_iso, "deleted_at": r.get("deleted_at") or None})
            stats["vehicles"] += 1

    # ---- Jasa
    for r in _rows(wb, "Jasa"):
        code, name = _s(r.get("code")), _s(r.get("name"))
        if not code or not name:
            continue
        data = {"name": name, "price": _i(r.get("price")), "active": _s(r.get("active")).lower() not in ("false", "0", "tidak", "")}
        existing = await db.services.find_one({"code": code})
        if existing:
            await db.services.update_one({"id": existing["id"]}, {"$set": {**data, "deleted_at": r.get("deleted_at") or None}})
        else:
            await db.services.insert_one({"id": new_id(), "code": code, **data, "created_at": now_iso, "deleted_at": r.get("deleted_at") or None})
            stats["services"] += 1

    # ---- Part (bulk)
    prows = _rows(wb, "Part")
    if prows:
        existing_by_code = {p["code"]: p for p in await db.parts.find({}, {"_id": 0}).to_list(None)}
        ops = []
        for r in prows:
            code, name = _s(r.get("code")), _s(r.get("name"))
            if not code or not name:
                continue
            data = {"name": name, "barcode": _s(r.get("barcode")), "price": _i(r.get("price")), "cost": _i(r.get("cost")),
                    "stock": _i(r.get("stock")), "min_stock": _i(r.get("min_stock")) or 5, "unit": _s(r.get("unit")) or "pcs",
                    "rack": _s(r.get("rack")), "active": _s(r.get("active")).lower() not in ("false", "0", "tidak", ""), "deleted_at": r.get("deleted_at") or None}
            ex = existing_by_code.get(code)
            if ex:
                ops.append(UpdateOne({"id": ex["id"]}, {"$set": data}))
            else:
                ops.append(InsertOne({"id": new_id(), "code": code, **data, "created_at": now_iso}))
                stats["parts"] += 1
        for i in range(0, len(ops), 1000):
            await db.parts.bulk_write(ops[i:i + 1000], ordered=False)

    # ---- Belanja
    for r in _rows(wb, "Belanja"):
        grp = _s(r.get("group")) or "BENGKEL"
        cat = _s(r.get("category"))
        if not cat:
            continue
        await db.expenses.insert_one({"id": new_id(), "_import": True, "group": grp, "category": cat, "amount": _i(r.get("amount")),
                                      "flow": _s(r.get("flow")) or "OUT", "note": _s(r.get("note")), "counterparty": _s(r.get("counterparty")),
                                      "supplier": _s(r.get("supplier")), "discount": _i(r.get("discount")), "het": _i(r.get("het")),
                                      "ongkir": _i(r.get("ongkir")), "unit_price": _i(r.get("unit_price")), "payment_method": _s(r.get("payment_method")) or "CASH",
                                      "part_name": _s(r.get("part_name")) or None, "qty": _i(r.get("qty")), "date": _s(r.get("date")),
                                      "created_by": _s(r.get("created_by")) or user["name"], "created_at": _s(r.get("created_at")) or now_iso, "deleted_at": r.get("deleted_at") or None})
        stats["expenses"] += 1

    # ---- Transaksi (join item/pembayaran by transaction_id; enrich metadata via created_at proximity)
    titems = _rows(wb, "Item Transaksi")
    tpays = _rows(wb, "Pembayaran")
    tcicilan = _rows(wb, "Cicilan Hutang")
    trx_rows = _rows(wb, "Transaksi")
    trx_sorted = sorted([t for t in trx_rows if _dt_of(t.get("created_at"))], key=lambda t: _dt_of(t.get("created_at")))
    # earliest item time per old txn id
    earliest: dict[str, datetime] = {}
    for it in titems:
        oid = _s(it.get("transaction_id"))
        d = _dt_of(it.get("created_at"))
        if oid and d and (oid not in earliest or d < earliest[oid]):
            earliest[oid] = d
    for p in tpays:
        oid = _s(p.get("transaction_id"))
        d = _dt_of(p.get("paid_at"))
        if oid and d and oid not in earliest:
            earliest[oid] = d

    def match_meta(oid: str) -> dict:
        # gunakan id langsung bila sheet Transaksi punya kolom id (backup versi baru)
        for t in trx_rows:
            if _s(t.get("id")) == oid:
                return t
        base = earliest.get(oid)
        if base is None or not trx_sorted:
            return {}
        best, bestdiff = None, None
        for t in trx_sorted:
            diff = abs((base - _dt_of(t.get("created_at"))).total_seconds())
            if bestdiff is None or diff < bestdiff:
                bestdiff, best = diff, t
        return best if (bestdiff is not None and bestdiff < 600) else {}

    items_by_oid: dict[str, list] = {}
    for it in titems:
        items_by_oid.setdefault(_s(it.get("transaction_id")), []).append(it)

    for oid, its in items_by_oid.items():
        if not oid:
            continue
        meta = match_meta(oid)
        date_ymd = _s(meta.get("date"))
        if not date_ymd:
            d0 = earliest.get(oid)
            date_ymd = d0.astimezone(timezone(timedelta(hours=7))).strftime("%Y%m%d") if d0 else wib_today()
        status = _s(meta.get("status")) or "SELESAI"
        trx_doc = {"id": oid, "_import": True, "trx_no": _s(meta.get("trx_no")) or f"TRX-{date_ymd}-{oid[:4]}",
                   "invoice_no": _s(meta.get("invoice_no")) or None, "queue_no": _s(meta.get("queue_no")), "date": date_ymd, "status": status,
                   "customer_name": _s(meta.get("customer_name")), "customer_phone": _s(meta.get("customer_phone")),
                   "plate": _s(meta.get("plate")), "vehicle_name": _s(meta.get("vehicle_name")), "km_in": _i(meta.get("km_in")),
                   "complaint_main": _s(meta.get("complaint_main")), "mechanic_name": _s(meta.get("mechanic_name")) or None,
                   "next_recommendation": _s(meta.get("next_recommendation")), "discount": _i(meta.get("discount")),
                   "debt_amount": _i(meta.get("debt_amount")), "debt_status": _s(meta.get("debt_status")) or None,
                   "debt_due_date": _s(meta.get("debt_due_date")), "paid_at": _s(meta.get("paid_at")) or None,
                   "created_at": _s(meta.get("created_at")) or now_iso, "updated_at": now_iso, "cancel_reason": _s(meta.get("cancel_reason")), "deleted_at": None}
        await db.service_transactions.insert_one(trx_doc)
        stats["transactions"] += 1
        for it in its:
            await db.service_items.insert_one({"id": new_id(), "_import": True, "transaction_id": oid, "kind": _s(it.get("kind")) or "part",
                                               "ref_id": None, "source": _s(it.get("source")) or "ESTIMASI", "approval": _s(it.get("approval")) or "DISETUJUI",
                                               "code": _s(it.get("code")), "name": _s(it.get("name")), "price": _i(it.get("price")), "cost": _i(it.get("cost")),
                                               "qty": _i(it.get("qty")) or 1, "subtotal": _i(it.get("subtotal")), "note": _s(it.get("note")),
                                               "added_by": _s(it.get("added_by")), "created_at": _s(it.get("created_at")) or now_iso, "deleted_at": it.get("deleted_at") or None})

    for p in tpays:
        oid = _s(p.get("transaction_id"))
        if not oid:
            continue
        await db.payments.insert_one({"id": new_id(), "_import": True, "transaction_id": oid, "method": _s(p.get("method")) or "CASH",
                                      "total": _i(p.get("total")), "amount_paid": _i(p.get("amount_paid")), "change": _i(p.get("change")),
                                      "discount": _i(p.get("discount")), "reference": _s(p.get("reference")), "cashier": _s(p.get("cashier")),
                                      "paid_at": _s(p.get("paid_at")) or now_iso, "date": _s(p.get("date")) or wib_today(), "due_date": _s(p.get("due_date")), "installments": []})
        stats["payments"] += 1
        # service_history agar tampil di histori & pengingat
        trx = await db.service_transactions.find_one({"id": oid}, {"_id": 0})
        if trx:
            await db.service_history.update_one({"transaction_id": oid}, {"$set": {"id": new_id(), "_import": True, "transaction_id": oid,
                                                "plate": trx.get("plate"), "invoice_no": trx.get("invoice_no"), "total": _i(p.get("total")),
                                                "mechanic_name": trx.get("mechanic_name"), "date": trx.get("date"), "created_at": _s(p.get("paid_at")) or now_iso}}, upsert=True)

    for c in tcicilan:
        oid = _s(c.get("transaction_id"))
        if not oid:
            continue
        await db.debt_payments.insert_one({"id": new_id(), "_import": True, "transaction_id": oid, "amount": _i(c.get("amount")),
                                           "method": _s(c.get("method")) or "CASH", "cashier": _s(c.get("cashier")), "date": _s(c.get("date")),
                                           "debt_before": _i(c.get("debt_before")), "debt_after": _i(c.get("debt_after")), "paid_at": _s(c.get("paid_at")) or now_iso})

    # ---- Penjualan Langsung (jika ada di backup versi baru)
    for r in _rows(wb, "Penjualan Langsung"):
        sale_no = _s(r.get("sale_no"))
        if not sale_no:
            continue
        await db.sales.insert_one({"id": new_id(), "_import": True, "sale_no": sale_no, "outlet_id": None, "outlet_name": _s(r.get("outlet_name")),
                                   "customer_name": _s(r.get("customer_name")) or "Umum", "customer_phone": _s(r.get("customer_phone")), "customer_address": _s(r.get("customer_address")),
                                   "items": [], "subtotal": _i(r.get("subtotal")), "discount": _i(r.get("discount")), "total": _i(r.get("total")),
                                   "total_cost": _i(r.get("total_cost")), "total_profit": _i(r.get("total_profit")), "method": _s(r.get("method")) or "CASH",
                                   "amount_paid": _i(r.get("amount_paid")), "change": 0, "debt_amount": _i(r.get("debt_amount")), "debt_status": _s(r.get("debt_status")) or None,
                                   "debt_due_date": "", "installments": [], "note": "", "date": _s(r.get("date")) or wib_today(), "cashier": _s(r.get("cashier")),
                                   "status": "SELESAI", "created_at": _s(r.get("created_at")) or now_iso, "deleted_at": r.get("deleted_at") or None})
        stats["sales"] += 1

    # ---- lanjutkan penomoran counter agar tidak bentrok
    async def bump_counter(key: str, coll, field: str, pattern: str):
        mx = 0
        async for d in db[coll].find({field: {"$regex": pattern}}, {field: 1, "_id": 0}):
            m = re.search(r"(\d+)(?!.*\d)", _s(d.get(field)))
            if m:
                mx = max(mx, int(m.group(1)))
        if mx:
            await db.counters.update_one({"_id": key}, {"$max": {"seq": mx}}, upsert=True)

    await bump_counter("trx", "service_transactions", "trx_no", r"^TRX-")
    await bump_counter("invoice", "service_transactions", "invoice_no", r"^NS-")
    await bump_counter("customer", "customers", "code", r"^C\d+")
    await bump_counter("sale", "sales", "sale_no", r"^PJ-")

    await audit(None, "BACKUP_IMPORT", user, f"Restore backup: {stats}")
    return {"ok": True, **stats}


# =========================================================================== #
#  PERBAIKAN BUG + FITUR BARU (cleanup part, keuangan, penggajian, SOP, hak akses)
# =========================================================================== #

def is_junk_part(code: str, name: str) -> bool:
    c = (code or "").strip().upper()
    n = (name or "").strip().lower()
    return bool(c) and (c[0] in ("Z", "L") or "kosong" in n or "low" in n)


async def sop_required_names() -> list[str]:
    items = await db.service_checklist_items.find({"active": True}, {"_id": 0}).sort("seq", 1).to_list(100)
    return [i["name"] for i in items]


WIB = timezone(timedelta(hours=7))


def month_default(month: str) -> str:
    return month or now().astimezone(WIB).strftime("%Y-%m")


# ------------------------------------------------------------------ Bug 2: cleanup part sampah & sampah import

@api.post("/parts/cleanup")
async def parts_cleanup(user: OwnerUser):
    """Pindahkan part sampah (Z*/L*/Kosong/Low) ke Sampah Import + deduplikasi kode."""
    parts = await db.parts.find({"deleted_at": None}, {"_id": 0}).to_list(None)
    moved = 0
    now_iso = iso(now())
    seen: dict[str, dict] = {}
    dup_removed = 0
    for p in parts:
        if is_junk_part(p.get("code", ""), p.get("name", "")):
            await db.import_trash.insert_one({"id": new_id(), "code": p.get("code"), "name": p.get("name"), "barcode": p.get("barcode", ""),
                                              "price": p.get("price", 0), "cost": p.get("cost", 0), "stock": p.get("stock", 0),
                                              "min_stock": p.get("min_stock", 5), "unit": p.get("unit", "pcs"), "rack": p.get("rack", ""),
                                              "reason": "Part sampah (kode Z/L atau nama Kosong/Low)", "source": "cleanup",
                                              "created_by": user["username"], "created_at": now_iso})
            await db.parts.update_one({"id": p["id"]}, {"$set": {"deleted_at": now_iso}})
            moved += 1
            continue
        code = (p.get("code") or "").strip().upper()
        if code in seen:
            # duplikat kode -> gabungkan stok ke yang pertama, sisanya di-soft delete
            keep = seen[code]
            await db.parts.update_one({"id": keep["id"]}, {"$inc": {"stock": int(p.get("stock", 0) or 0)}})
            keep["stock"] = int(keep.get("stock", 0) or 0) + int(p.get("stock", 0) or 0)
            await db.parts.update_one({"id": p["id"]}, {"$set": {"deleted_at": now_iso}})
            dup_removed += 1
        else:
            seen[code] = p
    await audit(None, "PARTS_CLEANUP", user, f"Bersihkan part: {moved} sampah dipindah, {dup_removed} duplikat digabung")
    return {"moved_to_trash": moved, "duplicates_merged": dup_removed, "remaining": len(seen)}


@api.get("/import-trash")
async def list_import_trash(_: OwnerUser):
    return await db.import_trash.find({"restored_at": None}, {"_id": 0}).sort("created_at", -1).to_list(2000)


@api.post("/import-trash/{tid}/restore")
async def restore_import_trash(tid: str, user: OwnerUser):
    t = await db.import_trash.find_one({"id": tid, "restored_at": None})
    if not t:
        raise HTTPException(status_code=404, detail="Data tidak ditemukan")
    code = (t.get("code") or "").strip()
    if await db.parts.find_one({"code": code, "deleted_at": None}):
        raise HTTPException(status_code=400, detail="Kode part sudah ada di master")
    doc = {"id": new_id(), "code": code, "name": t.get("name", ""), "barcode": t.get("barcode", ""), "price": int(t.get("price", 0) or 0),
           "cost": int(t.get("cost", 0) or 0), "stock": int(t.get("stock", 0) or 0), "min_stock": int(t.get("min_stock", 5) or 5),
           "unit": t.get("unit", "pcs"), "rack": t.get("rack", ""), "active": True, "created_at": iso(now()), "deleted_at": None}
    await db.parts.insert_one(doc)
    await db.import_trash.update_one({"id": tid}, {"$set": {"restored_at": iso(now())}})
    await audit(None, "IMPORT_TRASH_RESTORE", user, f"Pulihkan part {code} dari sampah import")
    return clean(doc)


@api.delete("/import-trash/{tid}")
async def delete_import_trash(tid: str, user: OwnerUser):
    r = await db.import_trash.delete_one({"id": tid})
    if not r.deleted_count:
        raise HTTPException(status_code=404, detail="Data tidak ditemukan")
    return {"ok": True}


# ------------------------------------------------------------------ Bug 4: Laporan Laba Rugi Harian & Kas

async def _modal_by_payment(pays: list[dict]) -> dict[str, int]:
    """Modal part (cost_snapshot) per transaction_id dari item servis."""
    tids = [p["transaction_id"] for p in pays]
    items = await db.service_items.find({"transaction_id": {"$in": tids}, "kind": "part", "deleted_at": None, "approval": "DISETUJUI"},
                                        {"_id": 0, "transaction_id": 1, "cost": 1, "qty": 1}).to_list(100000)
    out: dict[str, int] = {}
    for i in items:
        out[i["transaction_id"]] = out.get(i["transaction_id"], 0) + int(i.get("cost", 0) or 0) * int(i.get("qty", 0) or 0)
    return out


def pay_ymd(p: dict) -> str:
    d = str(p.get("date") or "")
    if len(d) == 8 and d.isdigit():
        return f"{d[:4]}-{d[4:6]}-{d[6:8]}"
    try:
        return datetime.fromisoformat(p["paid_at"]).astimezone(WIB).strftime("%Y-%m-%d")
    except (ValueError, TypeError, KeyError):
        return ""


async def build_laba_rugi(month: str, dfrom: str = "", dto: str = "") -> dict:
    pays = await db.payments.find({"cancelled": {"$ne": True}, "sale": {"$ne": True}, "date": ymd_range_filter(month, dfrom, dto)}, {"_id": 0}).to_list(50000)
    modal_map = await _modal_by_payment(pays)
    sales = await db.sales.find({"deleted_at": None, "date": ymd_range_filter(month, dfrom, dto)}, {"_id": 0}).to_list(50000)
    exp = await db.expenses.find({"deleted_at": None, "date": dash_range_filter(month, dfrom, dto)}, {"_id": 0}).to_list(50000)

    days: dict[str, dict] = {}

    def day(k: str) -> dict:
        return days.setdefault(k, {"date": k, "omzet": 0, "modal": 0, "belanja_bengkel": 0, "prive": 0, "laba_bersih": 0})

    for p in pays:
        k = pay_ymd(p)
        if not k:
            continue
        b = day(k)
        b["omzet"] += int(p.get("total", 0) or 0)
        b["modal"] += modal_map.get(p["transaction_id"], 0)
    for s in sales:
        d = str(s.get("date") or "")
        k = f"{d[:4]}-{d[4:6]}-{d[6:8]}" if len(d) == 8 else d
        if not k:
            continue
        b = day(k)
        b["omzet"] += int(s.get("total", 0) or 0)
        b["modal"] += int(s.get("total_cost", 0) or 0)
    for e in exp:
        k = str(e.get("date") or "")[:10]
        if not k:
            continue
        b = day(k)
        amt = int(e.get("amount", 0) or 0)
        if e.get("group") == "BENGKEL":
            b["belanja_bengkel"] += amt if e.get("flow", "OUT") == "OUT" else -amt
        elif e.get("group") == "KELUARGA":
            b["prive"] += amt
    for b in days.values():
        b["laba_bersih"] = b["omzet"] - b["modal"] - b["belanja_bengkel"]
    rows = sorted(days.values(), key=lambda r: r["date"], reverse=True)

    total = {"omzet": sum(r["omzet"] for r in rows), "modal": sum(r["modal"] for r in rows),
             "belanja_bengkel": sum(r["belanja_bengkel"] for r in rows), "prive": sum(r["prive"] for r in rows)}
    total["laba_bersih"] = total["omzet"] - total["modal"] - total["belanja_bengkel"]

    # Gaji mekanik (dari data karyawan aktif berjabatan mekanik) & rasio terhadap laba bersih
    emps = await db.employees.find({"active": True, "deleted_at": None}, {"_id": 0}).to_list(200)
    gaji_mekanik = sum(int(e.get("base_salary", 0) or 0) for e in emps if e.get("position") == "mekanik")
    pct = round((gaji_mekanik / total["laba_bersih"]) * 100, 1) if total["laba_bersih"] > 0 else None
    ideal = pct is not None and pct <= 30
    return {"month": month, "rows": rows, "total": total, "gaji_mekanik": gaji_mekanik,
            "gaji_pct": pct, "gaji_ideal": ideal}


@api.get("/reports/laba-rugi")
async def report_laba_rugi(_: OwnerUser, month: str = "", date_from: str = "", date_to: str = ""):
    return await build_laba_rugi(month_default(month), date_from, date_to)


@api.get("/reports/cashflow")
async def report_cashflow(_: OwnerUser, month: str = "", date_from: str = "", date_to: str = ""):
    month = month_default(month)
    fin = await db.settings.find_one({"id": "finance"}, {"_id": 0}) or {}
    awal_bengkel = int(fin.get("saldo_awal_bengkel", 0) or 0)
    awal_pribadi = int(fin.get("saldo_awal_pribadi", 0) or 0)
    pays = await db.payments.find({"cancelled": {"$ne": True}, "date": ymd_range_filter(month, date_from, date_to)}, {"_id": 0}).to_list(50000)
    exp = await db.expenses.find({"deleted_at": None, "date": dash_range_filter(month, date_from, date_to)}, {"_id": 0}).to_list(50000)
    # Pemasukan bengkel: pembayaran non-KREDIT (uang benar-benar masuk) + cicilan hutang + belanja BENGKEL flow IN
    masuk_tunai = sum(int(p.get("total", 0) or 0) for p in pays if p.get("method") != "KREDIT")
    cicilan = 0
    for p in pays:
        for ins in p.get("installments", []) or []:
            cicilan += int(ins.get("amount", 0) or 0)
    belanja_in = sum(int(e.get("amount", 0) or 0) for e in exp if e.get("group") == "BENGKEL" and e.get("flow") == "IN")
    belanja_out = sum(int(e.get("amount", 0) or 0) for e in exp if e.get("group") == "BENGKEL" and e.get("flow", "OUT") == "OUT")
    prive = sum(int(e.get("amount", 0) or 0) for e in exp if e.get("group") == "KELUARGA")
    pemasukan = masuk_tunai + cicilan + belanja_in
    saldo_bengkel = awal_bengkel + pemasukan - belanja_out - prive
    saldo_pribadi = awal_pribadi + prive  # dana yang diambil owner masuk ke kas pribadi
    return {
        "month": month,
        "bengkel": {"saldo_awal": awal_bengkel, "pemasukan": pemasukan, "pembayaran_tunai": masuk_tunai, "cicilan": cicilan,
                    "belanja_masuk": belanja_in, "pengeluaran": belanja_out, "prive": prive, "saldo_akhir": saldo_bengkel},
        "pribadi": {"saldo_awal": awal_pribadi, "prive_masuk": prive, "saldo_akhir": saldo_pribadi},
    }


# ------------------------------------------------------------------ Pengaturan Keuangan (Owner)

class FinanceIn(BaseModel):
    saldo_awal_bengkel: int = 0
    saldo_awal_pribadi: int = 0
    owner_draw: int = 0
    target_laba_bulanan: int = 0


@api.get("/settings/finance")
async def get_finance(_: OwnerUser):
    doc = await db.settings.find_one({"id": "finance"}, {"_id": 0}) or {}
    return {"saldo_awal_bengkel": int(doc.get("saldo_awal_bengkel", 0) or 0),
            "saldo_awal_pribadi": int(doc.get("saldo_awal_pribadi", 0) or 0),
            "owner_draw": int(doc.get("owner_draw", 0) or 0),
            "target_laba_bulanan": int(doc.get("target_laba_bulanan", 0) or 0)}


@api.put("/settings/finance")
async def set_finance(body: FinanceIn, user: OwnerUser):
    await db.settings.update_one({"id": "finance"}, {"$set": {"id": "finance", **body.model_dump()}}, upsert=True)
    await audit(None, "FINANCE_SETTINGS", user, "Ubah saldo awal / pengambilan owner")
    return body.model_dump()


# ------------------------------------------------------------------ Hak Akses per Role (Owner)

class AccessIn(BaseModel):
    roles: dict[str, dict[str, bool]]


@api.get("/settings/access")
async def get_access(_: CurrentUser):
    doc = await db.settings.find_one({"id": "access"}, {"_id": 0}) or {}
    return {"roles": doc.get("roles", {})}


@api.put("/settings/access")
async def set_access(body: AccessIn, user: OwnerUser):
    await db.settings.update_one({"id": "access"}, {"$set": {"id": "access", "roles": body.roles}}, upsert=True)
    await audit(None, "ACCESS_SETTINGS", user, "Ubah hak akses per role")
    return {"roles": body.roles}


# ------------------------------------------------------------------ Bug 6: SOP Checklist Items (Owner) + per transaksi

class ChecklistItemMaster(BaseModel):
    name: str
    active: bool = True


@api.get("/service-checklist-items")
async def list_sop_items(_: CurrentUser):
    return await db.service_checklist_items.find({"active": True}, {"_id": 0}).sort("seq", 1).to_list(100)


@api.post("/service-checklist-items")
async def create_sop_item(body: ChecklistItemMaster, user: OwnerUser):
    seq = await db.service_checklist_items.count_documents({})
    doc = {"id": new_id(), "name": body.name.strip(), "seq": seq, "active": True, "created_at": iso(now())}
    await db.service_checklist_items.insert_one(doc)
    return clean(doc)


@api.put("/service-checklist-items/{iid}")
async def update_sop_item(iid: str, body: ChecklistItemMaster, user: OwnerUser):
    doc = await db.service_checklist_items.find_one_and_update({"id": iid}, {"$set": {"name": body.name.strip(), "active": body.active}}, return_document=True)
    if not doc:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    return clean(doc)


@api.delete("/service-checklist-items/{iid}")
async def delete_sop_item(iid: str, user: OwnerUser):
    r = await db.service_checklist_items.update_one({"id": iid}, {"$set": {"active": False}})
    if not r.modified_count:
        raise HTTPException(status_code=404, detail="Item tidak ditemukan")
    return {"ok": True}


class SopSaveIn(BaseModel):
    items: list[dict]  # [{name, ok, note}]


@api.post("/transactions/{tid}/sop")
async def save_sop(tid: str, body: SopSaveIn, user: MekanikUser):
    trx = await get_trx(tid)
    if trx["status"] == "DIBATALKAN":
        raise HTTPException(status_code=400, detail="Transaksi sudah dibatalkan")
    req = await sop_required_names()
    done = {str(c.get("name")): bool(c.get("ok")) for c in body.items}
    sop_ok = all(done.get(n) for n in req) if req else True
    checklist = [{"name": c.get("name"), "ok": bool(c.get("ok")), "note": c.get("note", "")} for c in body.items]
    await db.service_transactions.update_one({"id": tid}, {"$set": {"sop_checklist": checklist, "sop_ok": sop_ok, "sop_by": user["name"], "sop_at": iso(now())}})
    await audit(tid, "SOP_CHECKLIST", user, f"Checklist SOP {'LENGKAP' if sop_ok else 'BELUM lengkap'}")
    return await build_detail(await get_trx(tid))


# ------------------------------------------------------------------ Penggajian (Payroll) — Owner only

class EmployeeIn(BaseModel):
    name: str
    position: Literal["mekanik", "kasir", "partman", "owner"]
    base_salary: int = 0
    bonus_per_unit: int = 0
    active: bool = True


@api.get("/employees")
async def list_employees(_: OwnerUser):
    return await db.employees.find({"deleted_at": None}, {"_id": 0}).sort("created_at", 1).to_list(500)


@api.post("/employees")
async def create_employee(body: EmployeeIn, user: OwnerUser):
    doc = {"id": new_id(), **body.model_dump(), "created_at": iso(now()), "deleted_at": None}
    await db.employees.insert_one(doc)
    await audit(None, "EMPLOYEE_CREATE", user, f"Tambah karyawan {body.name} ({body.position})")
    return clean(doc)


@api.put("/employees/{eid}")
async def update_employee(eid: str, body: EmployeeIn, user: OwnerUser):
    doc = await db.employees.find_one_and_update({"id": eid, "deleted_at": None}, {"$set": body.model_dump()}, return_document=True)
    if not doc:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan")
    await audit(None, "EMPLOYEE_UPDATE", user, f"Ubah karyawan {body.name}")
    return clean(doc)


@api.delete("/employees/{eid}")
async def delete_employee(eid: str, user: OwnerUser):
    r = await db.employees.update_one({"id": eid, "deleted_at": None}, {"$set": {"deleted_at": iso(now()), "active": False}})
    if not r.modified_count:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan")
    return {"ok": True}


async def mechanic_units(name: str, month: str) -> int:
    """Jumlah unit servis selesai (non-batal) atas nama mekanik pada bulan berjalan."""
    return await db.service_transactions.count_documents(
        {"mechanic_name": name, "status": {"$in": STATUS_PAID}, "deleted_at": None, "date": {"$regex": f"^{month.replace('-', '')}"}})


async def payslip(emp: dict, month: str) -> dict:
    base = int(emp.get("base_salary", 0) or 0)
    units = 0
    bonus = 0
    if emp.get("position") == "mekanik":
        units = await mechanic_units(emp["name"], month)
        bonus = units * int(emp.get("bonus_per_unit", 0) or 0)
    return {"employee_id": emp["id"], "name": emp["name"], "position": emp["position"], "period": month,
            "base_salary": base, "units": units, "bonus_per_unit": int(emp.get("bonus_per_unit", 0) or 0),
            "bonus": bonus, "total": base + bonus}


@api.get("/payroll/slip/{eid}")
async def payroll_slip(eid: str, _: OwnerUser, month: str = ""):
    emp = await db.employees.find_one({"id": eid, "deleted_at": None}, {"_id": 0})
    if not emp:
        raise HTTPException(status_code=404, detail="Karyawan tidak ditemukan")
    return await payslip(emp, month_default(month))


@api.get("/payroll/report")
async def payroll_report(_: OwnerUser, month: str = ""):
    month = month_default(month)
    emps = await db.employees.find({"deleted_at": None, "active": True}, {"_id": 0}).sort("position", 1).to_list(500)
    slips = [await payslip(e, month) for e in emps]
    fin = await db.settings.find_one({"id": "finance"}, {"_id": 0}) or {}
    owner_draw = int(fin.get("owner_draw", 0) or 0)
    return {"month": month, "slips": slips, "owner_draw": owner_draw,
            "total_gaji": sum(s["total"] for s in slips), "grand_total": sum(s["total"] for s in slips) + owner_draw,
            "count": len(slips)}


# ------------------------------------------------------------------ Bug 5: Pengingat Hutang (3 hari sebelum jatuh tempo)

@api.get("/debt-reminders")
async def debt_reminders(_: CurrentUser):
    today = now().astimezone(WIB).date()
    out = []
    async for t in db.service_transactions.find({"debt_status": "BELUM LUNAS", "deleted_at": None}, {"_id": 0}):
        due = parse_date(t.get("debt_due_date") or "")
        if not due:
            continue
        days = (due - today).days
        if days <= 3:
            out.append({"id": t["id"], "sale": False, "invoice_no": t.get("invoice_no"), "customer_name": t.get("customer_name"),
                        "plate": t.get("plate"), "debt_amount": int(t.get("debt_amount", 0) or 0), "due_date": due.isoformat(),
                        "days_left": days, "status": "TERLAMBAT" if days < 0 else "JATUH TEMPO" if days == 0 else "SEGERA"})
    async for s in db.sales.find({"debt_status": "BELUM LUNAS", "deleted_at": None}, {"_id": 0}):
        due = parse_date(s.get("debt_due_date") or "")
        if not due:
            continue
        days = (due - today).days
        if days <= 3:
            out.append({"id": s["id"], "sale": True, "invoice_no": s.get("sale_no"), "customer_name": s.get("customer_name"),
                        "plate": "PENJUALAN PART", "debt_amount": int(s.get("debt_amount", 0) or 0), "due_date": due.isoformat(),
                        "days_left": days, "status": "TERLAMBAT" if days < 0 else "JATUH TEMPO" if days == 0 else "SEGERA"})
    out.sort(key=lambda r: r["due_date"])
    return out


app.include_router(api)
app.add_middleware(CORSMiddleware, allow_credentials=True, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])